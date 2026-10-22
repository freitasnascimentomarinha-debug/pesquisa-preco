"""Testes offline das listas de verificação, dos modelos novos e do layout de ofício."""
from datetime import date

import documento_oficial as D
import modelos_processo as MP

DADOS = {"om": "Centro de Teste", "cidade": "Rio de Janeiro", "uf": "RJ",
         "assinantes": [{"nome": "Fulano", "posto": "CT (IM)", "cargo": "Encarregado"}]}


def test_limites_2026():
    assert MP.LIMITES_DISPENSA["inciso_I"] == 130_984.20
    assert MP.LIMITES_DISPENSA["inciso_II"] == 65_492.11
    assert "12.807/2025" in MP.nota_limites()
    assert MP.brl(65492.11) == "R$ 65.492,11"


def test_checklist_sem_citacoes_erradas():
    tudo = " ".join(f for itens in MP.CHECKLIST.values() for _t, f, _a in itens)
    assert "art. 84, §3º" not in tudo and "art. 86, §1º" not in tudo
    for itens in MP.CHECKLIST.values():
        for texto, fundamento, atalho in itens:
            assert texto and fundamento
            assert atalho is None or atalho[0] in ("doc", "pag")
    diretas = MP.CHECKLIST["Dispensa eletrônica (art. 75, I e II)"]
    fund = " ".join(f for _t, f, _a in diretas)
    for inciso in ("art. 72, IV", "art. 72, VI", "art. 72, VII", "art. 91, §4º", "art. 53, §§4º e 5º"):
        assert inciso in fund, inciso


def test_resumo_ignora_nao_se_aplica():
    r = MP.resumo_checklist(["Presente", "Não se aplica", "Pendente", "Presente"])
    assert r == {"presentes": 2, "pendentes": 1, "nao_aplica": 1, "aplicaveis": 3, "percentual": 67}


def test_dias_uteis():
    assert MP.dias_uteis(date(2026, 10, 19), date(2026, 10, 22)) == 3     # seg -> qui
    assert MP.dias_uteis(date(2026, 10, 23), date(2026, 10, 27)) == 2     # sex -> ter (fim de semana)
    assert MP.dias_uteis(date(2026, 4, 2), date(2026, 4, 6)) == 1         # Sexta-feira Santa 03/04/2026


def test_secoes_obrigatorias_viram_lacuna_e_opcionais_justificativa():
    corpo = MP.corpo_secoes(MP.MODELOS[MP.DOC_ETP], [""] * 13)
    blocos = D.interpretar(corpo, "ESTUDO TÉCNICO PRELIMINAR")
    textos = [b["texto"] for b in blocos if b["tipo"] == "texto"]
    assert textos[0].startswith("______")                                   # I: obrigatória
    assert "art. 18, §2º" in textos[1]                                       # II: justificar


def test_documentos_novos_geram_pdf_e_docx():
    corpos = [
        ("ESTUDO TÉCNICO PRELIMINAR", MP.corpo_secoes(MP.MODELOS[MP.DOC_ETP], ["texto"] * 13), None),
        ("MAPA DE RISCOS", MP.corpo_mapa_riscos("Papel", MP.RISCOS_PADRAO), None),
        ("AVISO", MP.corpo_aviso({"objeto": "Papel", "valor": 1000.0}), None),
        ("LISTA", MP.corpo_checklist("Termo Aditivo", MP.CHECKLIST["Termo Aditivo"],
                                     ["Pendente"] * 13, [""] * 13, {}), None),
        ("OFÍCIO", MP.corpo_oficio_adesao({"ata": "1/2026"}), {"numero": "1", "do": "A", "ao": "B", "assunto": "C",
                                                              "referencias": ["Lei"], "copias": ["Arquivo"]}),
        ("OFÍCIO", MP.corpo_oficio_irp({}), {"externo": True, "destinatario": "Ao Senhor\nX", "vocativo": "Senhor,"}),
    ]
    for titulo, corpo, oficio in corpos:
        assert D.gerar_pdf(titulo, corpo, DADOS, oficio)[:4] == b"%PDF"
        assert D.gerar_docx(titulo, corpo, DADOS, oficio)[:2] == b"PK"


def test_oficio_sem_titulo_e_externo_com_fecho():
    import io
    from docx import Document
    of = D.gerar_docx("OFÍCIO", "1. Texto.", DADOS, {"numero": "7", "do": "Diretor", "ao": "Chefe", "assunto": "Teste"})
    texto = "\n".join(p.text for p in Document(io.BytesIO(of)).paragraphs)
    assert "Nº 7" in texto and "Do:\tDiretor" in texto and "OFÍCIO" not in texto
    ext = D.gerar_docx("OFÍCIO", "1. Texto.", DADOS, {"externo": True, "numero": "7/X-MB", "vocativo": "Senhor Diretor,"})
    texto = "\n".join(p.text for p in Document(io.BytesIO(ext)).paragraphs)
    assert "Ofício nº 7/X-MB" in texto and "Senhor Diretor," in texto and "Atenciosamente," in texto


def test_larguras_cabem_na_pagina():
    tab = [b for b in D.interpretar(MP.corpo_mapa_riscos("x", MP.RISCOS_PADRAO)) if b["tipo"] == "tabela"][0]
    larg = D._larguras_colunas(tab["linhas"], 165, 9)
    assert abs(sum(larg) - 165) < 0.01 and min(larg) > 6


if __name__ == "__main__":
    for nome, f in list(globals().items()):
        if nome.startswith("test_"):
            f()
            print("ok", nome)
