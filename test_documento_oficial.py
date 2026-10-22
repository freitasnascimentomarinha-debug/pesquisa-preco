"""Testes offline do gerador de documentos oficiais (rodar com: python -m pytest test_documento_oficial.py)."""
import io
from datetime import date

import documento_oficial as D

CORPO = """Claro! Segue o documento:

## TERMO DE REFERÊNCIA

### 1. Objeto
Aquisição de **papel** para o *COpAb*, sob nº __________ e __________.
- primeira alínea
  - subalínea
- segunda alínea

| Item | Descrição | Valor (R$) |
|---|---|---|
| 1 | Papel A4 | 250,00 |

---
Documento gerado pelo sistema O Babilaca (IA) — Ferramenta de apoio.
"""


def test_local_e_data():
    assert D.local_e_data("Rio de Janeiro", "rj", date(2026, 3, 1)) == "Rio de Janeiro, RJ, 1º de março de 2026."
    assert D.local_e_data("", "", date(2026, 3, 15)).endswith("15 de março de 2026.")


def test_interpretar_limpa_conversa_titulo_e_rodape():
    blocos = D.interpretar(CORPO, "TERMO DE REFERÊNCIA")
    tipos = [b["tipo"] for b in blocos]
    assert blocos[0] == {"tipo": "item", "texto": "1. OBJETO"}
    assert [b["texto"] for b in blocos if b["tipo"] in ("alinea", "subalinea")] == ["a) primeira alínea", "I) subalínea", "b) segunda alínea"]
    assert "tabela" in tipos
    assert not any("Claro" in b.get("texto", "") or "Babilaca" in b.get("texto", "") for b in blocos)


def test_lacunas_nao_viram_italico():
    assert D.segmentos("nº __________ e __________ fim") == [("nº __________ e __________ fim", False, False)]
    assert D.segmentos("a **b** *c*") == [("a ", False, False), ("b", True, False), (" ", False, False), ("c", False, True)]


def test_gera_pdf_e_docx():
    import docx
    pdf = D.gerar_pdf("TERMO DE REFERÊNCIA", CORPO, {"om": "Centro de Operações do Abastecimento"})
    assert pdf.startswith(b"%PDF")
    documento = docx.Document(io.BytesIO(D.gerar_docx("TERMO DE REFERÊNCIA", CORPO, {"om": "COpAb"})))
    textos = [p.text for p in documento.paragraphs]
    assert "MARINHA DO BRASIL" in textos and "COPAB" in textos and len(documento.tables) == 1
