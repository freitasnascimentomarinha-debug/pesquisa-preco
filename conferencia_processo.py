"""Conferência automática do processo de contratação a partir dos PDFs (e Word) enviados.

Tudo aqui é regra objetiva, sem IA: identifica as peças pelo título de cada página, marca a lista de verificação da
modalidade, e confere coerência (CNPJ, valores, datas), validade de certidões, limite de dispensa, prazo do aviso,
ordem das peças e lacunas não preenchidas. Cada achado diz onde está (arquivo e página) para o usuário conferir.

Limites conhecidos: PDF escaneado (sem texto) não é lido; a identificação depende de o título da peça estar no topo
da página; valores e datas são lidos por padrão de texto e podem falhar em layouts incomuns.
"""

from __future__ import annotations

import io
import re
import unicodedata
from collections import Counter
from datetime import date, datetime

import modelos_processo as MP

# ── Leitura ───────────────────────────────────────────────────────────────────────────────────────────────────


def extrair_paginas(arquivos: list[tuple[str, bytes]]) -> list[dict]:
    """[(nome, conteúdo)] -> [{'arquivo', 'pagina', 'texto'}], na ordem enviada. Word vira uma 'página' só."""
    paginas = []
    for nome, conteudo in arquivos:
        if nome.lower().endswith(".pdf"):
            try:
                try:
                    from PyPDF2 import PdfReader
                except ImportError:  # mesma API no sucessor pypdf
                    from pypdf import PdfReader
                leitor = PdfReader(io.BytesIO(conteudo))
                for n, pag in enumerate(leitor.pages, 1):
                    try:
                        texto = pag.extract_text() or ""
                    except Exception:
                        texto = ""
                    paginas.append({"arquivo": nome, "pagina": n, "texto": texto})
            except Exception as erro:
                paginas.append({"arquivo": nome, "pagina": 0, "texto": "", "erro": f"PDF ilegível ({type(erro).__name__})"})
        elif nome.lower().endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(conteudo))
            texto = "\n".join(p.text for p in doc.paragraphs)
            for tab in doc.tables:
                texto += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in tab.rows)
            paginas.append({"arquivo": nome, "pagina": 1, "texto": texto})
    return paginas


def _sem_acento(texto: str) -> str:
    return unicodedata.normalize("NFD", texto or "").encode("ascii", "ignore").decode().lower()


def onde(paginas: list[dict], indices: list[int]) -> str:
    """'arquivo.pdf, p. 3-5; outro.pdf, p. 1'."""
    grupos: dict[str, list[int]] = {}
    for i in indices:
        grupos.setdefault(paginas[i]["arquivo"], []).append(paginas[i]["pagina"])
    partes = []
    for arq, nums in grupos.items():
        nums = sorted(set(nums))
        faixas, ini = [], nums[0]
        for a, b in zip(nums, nums[1:] + [None]):
            if b != a + 1:
                faixas.append(str(ini) if ini == a else f"{ini}-{a}")
                ini = b
        partes.append(f"{arq}, p. {', '.join(faixas)}")
    return "; ".join(partes)


# ── Identificação das peças ──────────────────────────────────────────────────────────────────────────────────
# (tipo, rótulo, padrão procurado no topo da página, sem acento). A ordem importa: o primeiro que casar vence.
TIPOS = [
    ("parecer", "Parecer jurídico", r"\bparecer\b"),
    ("dfd", "DFD", r"formalizacao de demanda|\bdfd\b"),
    ("etp", "Estudo Técnico Preliminar", r"estudo tecnico preliminar"),
    ("riscos", "Mapa de Riscos", r"mapa de riscos|gerenciamento de riscos|analise de riscos"),
    ("tjdl", "Justificativa de dispensa (TJDL)", r"justificativa de dispensa|dispensa de licitacao"),
    ("tjil", "Justificativa de inexigibilidade (TJIL)", r"inexigibilidade"),
    ("aviso", "Aviso de contratação direta", r"aviso de (contratacao direta|dispensa)"),
    ("edital", "Edital", r"\bedital\b"),
    ("tr", "Termo de Referência", r"termo de referencia|projeto basico"),
    ("mapa", "Mapa comparativo / pesquisa de preços", r"mapa comparativo|pesquisa de precos|quadro comparativo|"
                                                      r"relatorio de (pesquisa|cotacao)"),
    ("proposta", "Proposta / orçamento do fornecedor", r"\bproposta\b|\borcamento\b|\bcotacao\b"),
    ("sicaf", "SICAF", r"sicaf|sistema de cadastramento unificado"),
    ("ceis", "Consulta CEIS/CNEP", r"\bceis\b|\bcnep\b|empresas inidoneas|consulta consolidada de pessoa juridica"),
    ("cadin", "Consulta CADIN", r"\bcadin\b"),
    ("certidao", "Certidão de regularidade", r"certidao|certificado de regularidade|\bcndt\b|\bcrf\b"),
    ("orcamento", "Disponibilidade orçamentária", r"nota de credito|disponibilidade orcamentaria|"
                                                 r"previsao (de recursos|orcamentaria)|declaracao orcamentaria"),
    ("oficio", "Ofício", r"\boficio\b"),
    ("ata_sessao", "Ata / relatório da sessão", r"ata da sessao|relatorio da (sessao|disputa)|termo de julgamento|"
                                               r"resultado da (sessao|dispensa)"),
    ("ata_rp", "Ata de Registro de Preços", r"ata de registro de precos"),
    ("homologacao", "Adjudicação e homologação", r"homologacao|adjudicacao|termo de homologacao"),
    ("autorizacao", "Autorização / ratificação", r"\bautorizo\b|\bratifico\b|autorizacao da (autoridade|despesa|contratacao)|"
                                                  r"despacho (decisorio|de autorizacao)"),
    ("empenho", "Nota de empenho", r"nota de empenho|\bempenho\b"),
    ("contrato", "Contrato", r"\bcontrato n|termo de contrato|instrumento contratual"),
    ("aditivo", "Termo aditivo", r"termo aditivo"),
    ("pncp", "Publicação no PNCP", r"portal nacional de contratacoes publicas|\bpncp\b"),
]
ROTULO = {t: r for t, r, _ in TIPOS}
TOPO = 450  # caracteres do topo da página onde se procura o título da peça


def classificar(paginas: list[dict]) -> list[str | None]:
    """Tipo de cada página. Página sem título reconhecível continua a peça anterior do mesmo arquivo."""
    tipos: list[str | None] = []
    for i, p in enumerate(paginas):
        topo = _sem_acento(p["texto"][:TOPO])
        achado = next((t for t, _r, padrao in TIPOS if re.search(padrao, topo)), None)
        if achado is None and i and paginas[i - 1]["arquivo"] == p["arquivo"]:
            achado = tipos[-1]
        tipos.append(achado)
    return tipos


def documentos(paginas: list[dict], tipos: list[str | None]) -> list[dict]:
    """Peças = páginas seguidas do mesmo tipo e arquivo: [{'tipo', 'indices', 'texto', 'data'}]."""
    docs: list[dict] = []
    for i, t in enumerate(tipos):
        if t is None:
            continue
        if docs and docs[-1]["tipo"] == t and docs[-1]["indices"][-1] == i - 1 \
                and paginas[i - 1]["arquivo"] == paginas[i]["arquivo"]:
            docs[-1]["indices"].append(i)
        else:
            docs.append({"tipo": t, "indices": [i]})
    for d in docs:
        d["texto"] = "\n".join(paginas[i]["texto"] for i in d["indices"])
        d["data"] = data_do_documento(d["texto"])
    return docs


# ── Extração de dados ────────────────────────────────────────────────────────────────────────────────────────
MESES = {m: n for n, m in enumerate(["janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho", "agosto",
                                     "setembro", "outubro", "novembro", "dezembro"], 1)}
RE_DATA_EXTENSO = re.compile(r"(\d{1,2})(?:º|°|o)?\s+de\s+([a-zç]+)\s+de\s+(\d{4})")
RE_DATA = re.compile(r"\b(\d{2})/(\d{2})/(\d{4})\b")
RE_CNPJ = re.compile(r"\b\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}\b")
RE_VALOR = re.compile(r"R\$\s*([\d.]+,\d{2})")


def _data(d: str, m: str | int, a: str) -> date | None:
    try:
        mes = MESES.get(_sem_acento(m)) if isinstance(m, str) and not m.isdigit() else int(m)
        return date(int(a), int(mes), int(d)) if mes else None
    except (ValueError, TypeError):
        return None


def datas(texto: str) -> list[date]:
    achadas = [_data(*m.groups()) for m in RE_DATA_EXTENSO.finditer(_sem_acento(texto))]
    achadas += [_data(*m.groups()) for m in RE_DATA.finditer(texto)]
    return [d for d in achadas if d and 2000 <= d.year <= 2100]


def data_do_documento(texto: str) -> date | None:
    """Data por extenso (local e data) se houver; senão a primeira dd/mm/aaaa."""
    m = RE_DATA_EXTENSO.search(_sem_acento(texto))
    if m and _data(*m.groups()):
        return _data(*m.groups())
    m = RE_DATA.search(texto)
    return _data(*m.groups()) if m else None


def valor_br(texto: str) -> float:
    return float(texto.replace(".", "").replace(",", "."))


def valores(texto: str) -> list[float]:
    return [valor_br(v) for v in RE_VALOR.findall(texto)]


def valor_total(texto: str) -> float | None:
    """Valor junto de 'total' / 'valor global' / 'valor estimado'; senão o maior valor do documento."""
    sa = _sem_acento(texto)
    perto = [valor_br(m.group(2)) for m in re.finditer(r"(total|valor global|valor estimado|valor da contratacao)"
                                                      r"[^\n]{0,80}?r\$\s*([\d.]+,\d{2})", sa)]
    if perto:
        return max(perto)
    todos = valores(texto)
    return max(todos) if todos else None


def cnpjs(texto: str) -> list[str]:
    from servicos_publicos import cnpj_valido, limpar_cnpj
    vistos = []
    for c in RE_CNPJ.findall(texto):
        d = limpar_cnpj(c)
        if len(d) == 14 and cnpj_valido(d) and d not in vistos:
            vistos.append(d)
    return vistos


def fc(cnpj: str) -> str:
    from servicos_publicos import formatar_cnpj
    return formatar_cnpj(cnpj)


def privado(cnpj: str) -> bool:
    """Órgãos federais costumam ter CNPJ iniciado por '00'; o fornecedor, não. Heurística para achar o contratado."""
    return not cnpj.startswith("00")


def validades(texto: str) -> list[date]:
    """Datas de validade de certidões: 'Válida até 10/05/2026', 'Validade: 01/05/2026 a 30/05/2026'."""
    achadas = []
    for m in re.finditer(r"valid[ao]\w*\s*(?:ate|:|-)?\s*(?:de\s*)?(\d{2}/\d{2}/\d{4})(?:\s*(?:a|ate)\s*(\d{2}/\d{2}/\d{4}))?",
                         _sem_acento(texto)):
        bruto = m.group(2) or m.group(1)
        d = _data(*bruto.split("/"))
        if d:
            achadas.append(d)
    return achadas


def fmt(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else "-"


def brl(v: float | None) -> str:
    return MP.brl(v) if v is not None else "-"


# ── Lista de verificação automática ──────────────────────────────────────────────────────────────────────────
# (padrão no texto do item, tipos de peça que o atendem, frases que o atendem em qualquer página)
REGRAS_ITENS = [
    (r"copia da ata", {"ata_rp", "edital"}, []),
    (r"formalizacao de demanda", {"dfd"}, []),
    (r"estudo tecnico preliminar", {"etp"}, []),
    (r"mapa de riscos|analise dos riscos", {"riscos"}, []),
    (r"termo de justificativa de dispensa", {"tjdl"}, []),
    (r"justificativa de inexigibilidade", {"tjil"}, []),
    (r"inviabilidade de competicao", set(), ["exclusividade", "notoria especializacao", "fornecedor exclusivo"]),
    (r"emergencia", set(), ["emergencia", "emergencial", "calamidade"]),
    (r"termo de referencia", {"tr"}, []),
    (r"mapa comparativo", {"mapa"}, ["mapa comparativo", "quadro comparativo"]),
    (r"pesquisa de precos|estimativa", {"mapa"}, ["pesquisa de precos", "preco estimado", "valor estimado"]),
    (r"recursos orcamentarios|creditos orcamentarios", {"orcamento"}, ["nota de credito", "disponibilidade orcamentaria",
                                                                      "dotacao orcamentaria"]),
    (r"aviso de contratacao", {"aviso"}, ["aviso de contratacao direta", "aviso de dispensa"]),
    (r"propostas e lances|ata da sessao", {"ata_sessao"}, []),
    (r"proposta formal", {"proposta"}, []),
    (r"\bceis\b", {"ceis"}, ["cadastro nacional de empresas inidoneas", "ceis", "cnep"]),
    (r"\bcadin\b", {"cadin"}, ["cadin"]),
    (r"habilitacao|regularidade", {"sicaf", "certidao"}, []),
    (r"razao da escolha", set(), ["razao da escolha", "escolha do fornecedor", "escolha do contratado"]),
    (r"justificativa de preco", set(), ["justificativa do preco", "justificativa de preco", "compatibilidade do preco",
                                        "compativel com os precos", "compativeis com os praticados"]),
    (r"parecer", {"parecer"}, []),
    (r"designacao", set(), ["pregoeiro", "agente de contratacao", "equipe de apoio"]),
    (r"adjudicacao e homologacao", {"homologacao"}, ["homologo", "adjudico"]),
    (r"ofic", {"oficio"}, []),
    (r"aceitacao do orgao gerenciador|manifestacao do orgao gerenciador", set(),
     ["autorizo a adesao", "autoriza a adesao", "aceite da adesao", "autorizacao para adesao", "de acordo com a adesao"]),
    (r"aceitacao do fornecedor|concordancia", set(), ["concordancia", "de acordo", "aceite", "manifesta interesse"]),
    (r"termo aditivo assinado|termo de prorrogacao", {"aditivo"}, ["termo de prorrogacao", "termo aditivo"]),
    (r"autorizacao", {"autorizacao"}, ["autorizo", "ratifico"]),
    (r"divulgacao|divulgad|pncp", {"pncp"}, ["pncp", "portal nacional de contratacoes"]),
    (r"contrato ou instrumento|nota de empenho|empenho", {"empenho", "contrato"}, ["nota de empenho"]),
    (r"edital", {"edital"}, []),
]


def checklist_automatico(modalidade: str, paginas: list[dict], docs: list[dict], total: float | None) -> list[dict]:
    """[{'texto', 'fundamento', 'situacao' ('Presente' | 'Não encontrado' | 'Verificar' | 'Não atende'), 'onde'}]."""
    texto_sa = [_sem_acento(p["texto"]) for p in paginas]
    saida = []
    for texto, fundamento, _a in MP.CHECKLIST[modalidade]:
        item_sa = _sem_acento(texto)
        situacao, local = "Verificar", "conferir manualmente"
        if "enquadramento no limite" in item_sa:
            limite = MP.LIMITES_DISPENSA["inciso_II"]
            if total is None:
                situacao, local = "Verificar", "valor total não identificado"
            else:
                situacao = ("Presente" if total < limite else
                            "Verificar" if total < MP.LIMITES_DISPENSA["inciso_I"] else "Não atende")
                local = f"total identificado {brl(total)} (limite do inciso II: {brl(limite)})"
        else:
            regra = next(((t, f) for padrao, t, f in REGRAS_ITENS if re.search(padrao, item_sa)), None)
            if regra:
                tipos_ok, frases = regra
                idx = [i for d in docs if d["tipo"] in tipos_ok for i in d["indices"]]
                if not idx and frases:
                    idx = [i for i, t in enumerate(texto_sa) if any(f in t for f in frases)]
                situacao = "Presente" if idx else "Não encontrado"
                local = onde(paginas, idx[:12]) if idx else "-"
        saida.append({"texto": texto, "fundamento": fundamento, "situacao": situacao, "onde": local})
    return saida


# ── Regras de conferência ────────────────────────────────────────────────────────────────────────────────────
GRAVE, CORRIGIR, SUGESTAO = "🔴", "🟡", "🔵"
ORDEM_GRAVIDADE = {GRAVE: 0, CORRIGIR: 1, SUGESTAO: 2}
LACUNAS = re.compile(r"_{5,}|\[\s*(preencher|inserir|xx+)[^\]]*\]|\bXX/XX/XXXX\b|\bR\$\s*X", re.I)


def conferir(paginas: list[dict], modalidade: str, data_referencia: date | None = None) -> dict:
    """Roda a conferência completa. Devolve {'docs', 'checklist', 'achados', 'dados', 'avisos'}."""
    tipos = classificar(paginas)
    docs = documentos(paginas, tipos)
    por_tipo: dict[str, list[dict]] = {}
    for d in docs:
        por_tipo.setdefault(d["tipo"], []).append(d)
    achados: list[dict] = []
    avisos: list[str] = []

    def achado(grav, titulo, detalhe, idx=None):
        achados.append({"gravidade": grav, "titulo": titulo, "detalhe": detalhe,
                        "onde": onde(paginas, idx) if idx else "-"})

    sem_texto = [i for i, p in enumerate(paginas) if len(p["texto"].strip()) < 20]
    if sem_texto:
        avisos.append(f"{len(sem_texto)} página(s) sem texto legível (provável digitalização/imagem): "
                      f"{onde(paginas, sem_texto[:15])}. Essas páginas não foram conferidas.")
    for p in paginas:
        if p.get("erro"):
            avisos.append(f"{p['arquivo']}: {p['erro']}.")

    # datas de referência
    def data_de(tipo):
        ds = [d["data"] for d in por_tipo.get(tipo, []) if d["data"]]
        return max(ds) if ds else None
    d_autorizacao, d_parecer, d_empenho = data_de("autorizacao"), data_de("parecer"), data_de("empenho")
    referencia = data_referencia or d_empenho or d_autorizacao or date.today()

    # 1. valores
    total_mapa = next((valor_total(d["texto"]) for d in por_tipo.get("mapa", []) if valor_total(d["texto"])), None)
    total_tr = next((valor_total(d["texto"]) for d in por_tipo.get("tr", []) + por_tipo.get("etp", [])
                     if valor_total(d["texto"])), None)
    total = total_mapa or total_tr
    if total_mapa and total_tr and abs(total_mapa - total_tr) > max(0.01 * total_mapa, 1):
        achado(CORRIGIR, "Valor estimado diferente entre o mapa de preços e o TR/ETP",
               f"Mapa: {brl(total_mapa)}; TR/ETP: {brl(total_tr)}. Confira qual é o valor correto e uniformize.",
               por_tipo["mapa"][0]["indices"][:1] + (por_tipo.get("tr") or por_tipo.get("etp"))[0]["indices"][:1])
    if total and "art. 75, I e II" in modalidade:
        if total >= MP.LIMITES_DISPENSA["inciso_I"]:
            achado(GRAVE, "Valor acima dos limites de dispensa", f"Total identificado {brl(total)}, acima de "
                   f"{brl(MP.LIMITES_DISPENSA['inciso_I'])} (art. 75, I). {MP.nota_limites()}")
        elif total >= MP.LIMITES_DISPENSA["inciso_II"]:
            achado(CORRIGIR, "Valor acima do limite do art. 75, II", f"Total identificado {brl(total)}, acima de "
                   f"{brl(MP.LIMITES_DISPENSA['inciso_II'])}: só cabe se o objeto for obra, serviço de engenharia ou "
                   "manutenção de veículos (art. 75, I). Considere também o somatório do exercício (art. 75, §1º).")
    v_credito = max((valor_total(d["texto"]) or 0 for d in por_tipo.get("orcamento", [])), default=0)
    if v_credito and total and v_credito + 0.01 < total:
        achado(GRAVE, "Crédito orçamentário menor que o valor estimado",
               f"Crédito identificado {brl(v_credito)} para total de {brl(total)} (art. 72, IV, e art. 150).",
               por_tipo["orcamento"][0]["indices"])
    v_empenho = max((valor_total(d["texto"]) or 0 for d in por_tipo.get("empenho", [])), default=0)
    if v_empenho and total and v_empenho > total * 1.001:
        achado(CORRIGIR, "Empenho maior que o valor estimado",
               f"Empenho {brl(v_empenho)} x estimado {brl(total)}: justifique a diferença.", por_tipo["empenho"][0]["indices"])

    # 2. fornecedor (CNPJ)
    cnpj_prop = Counter(c for d in por_tipo.get("proposta", []) + por_tipo.get("tjdl", []) + por_tipo.get("tjil", [])
                        for c in cnpjs(d["texto"]) if privado(c))
    cnpj_hab = Counter(c for t in ("certidao", "sicaf", "ceis", "cadin") for d in por_tipo.get(t, [])
                       for c in cnpjs(d["texto"]) if privado(c))
    contratado = (cnpj_prop or cnpj_hab).most_common(1)[0][0] if (cnpj_prop or cnpj_hab) else None
    if contratado and cnpj_hab and contratado not in cnpj_hab:
        outros = ", ".join(fc(c) for c in sorted(cnpj_hab))
        achado(GRAVE, "Certidões de CNPJ diferente do fornecedor escolhido",
               f"Fornecedor da proposta/justificativa: {fc(contratado)}; certidões/consultas: {outros}.",
               [i for t in ("certidao", "sicaf", "ceis") for d in por_tipo.get(t, []) for i in d["indices"]][:6])
    if contratado and por_tipo.get("empenho"):
        no_empenho = {c for d in por_tipo["empenho"] for c in cnpjs(d["texto"])}
        if no_empenho and contratado not in no_empenho:
            achado(GRAVE, "Nota de empenho em nome de outro CNPJ",
                   f"Esperado {fc(contratado)}; no empenho: {', '.join(fc(c) for c in sorted(no_empenho))}.", por_tipo["empenho"][0]["indices"])

    # 3. certidões
    for t in ("certidao", "sicaf"):
        for d in por_tipo.get(t, []):
            vals = validades(d["texto"])
            vencidas = [v for v in vals if v < referencia]
            if vencidas:
                achado(GRAVE, f"{ROTULO[t]} vencida na data de referência",
                       f"Validade {', '.join(fmt(v) for v in sorted(set(vencidas)))}; referência {fmt(referencia)} "
                       "(data do empenho, da autorização ou a informada). Emita nova certidão.", d["indices"])
            elif vals and min(vals) <= date.fromordinal(referencia.toordinal() + 5):
                achado(CORRIGIR, f"{ROTULO[t]} vence em até 5 dias", f"Validade {fmt(min(vals))}.", d["indices"])

    # 4. datas e ordem
    if d_parecer and d_autorizacao and d_autorizacao < d_parecer:
        achado(GRAVE, "Autorização anterior ao parecer jurídico",
               f"Autorização em {fmt(d_autorizacao)} e parecer em {fmt(d_parecer)}: a autorização deve se apoiar no parecer.",
               por_tipo["autorizacao"][0]["indices"][:1] + por_tipo["parecer"][0]["indices"][:1])
    if d_empenho and d_autorizacao and d_empenho < d_autorizacao:
        achado(GRAVE, "Empenho anterior à autorização da contratação",
               f"Empenho em {fmt(d_empenho)} e autorização em {fmt(d_autorizacao)} (art. 72, VIII).",
               por_tipo["empenho"][0]["indices"][:1])
    d_mapa = data_de("mapa")
    if d_mapa and (referencia - d_mapa).days > 180:
        achado(CORRIGIR, "Pesquisa de preços com mais de 6 meses",
               f"Pesquisa datada de {fmt(d_mapa)} ({(referencia - d_mapa).days} dias antes da referência). Confira os "
               "prazos de validade das fontes no art. 5º da IN SEGES/ME nº 65/2021 e atualize se preciso.",
               por_tipo["mapa"][0]["indices"][:1])
    for d in por_tipo.get("aviso", []):
        ds = sorted(set(datas(d["texto"])))
        if len(ds) >= 2:
            uteis = MP.dias_uteis(ds[0], ds[-1])
            if uteis < 3:
                achado(GRAVE, "Aviso de contratação direta com prazo menor que 3 dias úteis",
                       f"De {fmt(ds[0])} a {fmt(ds[-1])}: {uteis} dia(s) útil(eis) (art. 75, §3º).", d["indices"])
    sequencia = ["dfd", "tr", "mapa", "parecer", "autorizacao", "empenho"]
    primeiros = [(t, por_tipo[t][0]["indices"][0]) for t in sequencia if por_tipo.get(t)]
    fora = [(a, b) for (a, ia), (b, ib) in zip(primeiros, primeiros[1:]) if ib < ia]
    for a, b in fora:
        achado(SUGESTAO, "Peças fora da ordem usual do processo",
               f"{ROTULO[b]} aparece antes de {ROTULO[a]}. As peças são juntadas na ordem em que são produzidas.",
               [por_tipo[b][0]["indices"][0], por_tipo[a][0]["indices"][0]])

    # 5. lacunas e campos não preenchidos
    com_lacuna = [i for i, p in enumerate(paginas) if LACUNAS.search(p["texto"])]
    if com_lacuna:
        achado(CORRIGIR, "Campos não preenchidos (lacunas ou marcadores)",
               "Há trechos como '_____', '[preencher]' ou 'XX/XX/XXXX' nas páginas indicadas.", com_lacuna[:15])

    # 6. assinatura nas peças decisórias
    sem_assinatura = [d for t in ("dfd", "tr", "parecer", "autorizacao") for d in por_tipo.get(t, [])
                      if not re.search(r"assinad[oa]|assinatura|gov\.br|certificado digital|codigo verificador|"
                                       r"autenticidade|\bcrc\b", _sem_acento(d["texto"]))]
    if sem_assinatura:
        achado(SUGESTAO, "Peças sem indicação de assinatura no texto",
               ", ".join(ROTULO[d["tipo"]] for d in sem_assinatura) + ": não achei 'assinado digitalmente' nem código "
               "de autenticidade. Se a assinatura for imagem ou carimbo, desconsidere.",
               [d["indices"][0] for d in sem_assinatura])

    checklist = checklist_automatico(modalidade, paginas, docs, total)
    faltando = []
    for c in checklist:
        item_sa = _sem_acento(c["texto"])
        if c["situacao"] == "Não encontrado" and re.search(r"se for o caso|quando exigid|quando cabivel|facultativ|"
                                                          r"\(se nao for|se houver", item_sa):
            c["situacao"], c["onde"] = "Verificar", "não encontrado; item exigido só em alguns casos"
        elif c["situacao"] == "Não encontrado":
            faltando.append(c["texto"])
    if faltando:
        achado(CORRIGIR, f"Itens da lista de verificação não encontrados ({len(faltando)})", "; ".join(faltando) + ".")
    achados.sort(key=lambda a: ORDEM_GRAVIDADE[a["gravidade"]])
    nao_identificadas = sum(1 for t in tipos if t is None)
    if nao_identificadas:
        avisos.append(f"{nao_identificadas} página(s) sem peça identificada pelo título.")
    return {"docs": docs, "checklist": checklist, "achados": achados, "avisos": avisos, "paginas": paginas,
            "dados": {"total": total, "total_mapa": total_mapa, "total_tr": total_tr, "credito": v_credito or None,
                      "empenho": v_empenho or None, "contratado": contratado, "referencia": referencia}}


# ── Relatório ────────────────────────────────────────────────────────────────────────────────────────────────
def _c(t: str) -> str:
    return (t or "").replace("|", "/").replace("\n", " ").strip()


def corpo_relatorio(r: dict, modalidade: str, arquivos: list[str], externos: list[dict] | None = None) -> str:
    p = r["paginas"]
    dd = r["dados"]
    cont = Counter(a["gravidade"] for a in r["achados"])
    linhas = ["### 1. IDENTIFICAÇÃO",
              f"**Modalidade / tipo de processo:** {modalidade}.",
              f"**Arquivos conferidos:** {', '.join(arquivos)} ({len(p)} página(s)).",
              f"**Data de referência:** {fmt(dd['referencia'])}.",
              f"**Fornecedor identificado:** {fc(dd['contratado']) if dd['contratado'] else 'não identificado'}. "
              f"**Valor estimado identificado:** {brl(dd['total'])}.",
              "", "### 2. RESULTADO",
              f"{cont.get(GRAVE, 0)} inconsistência(s) grave(s), {cont.get(CORRIGIR, 0)} a corrigir e "
              f"{cont.get(SUGESTAO, 0)} sugestão(ões).",
              "", "| Gravidade | Achado | Detalhe | Onde |", "|---|---|---|---|"]
    nome = {GRAVE: "Grave", CORRIGIR: "Corrigir", SUGESTAO: "Sugestão"}
    for a in r["achados"]:
        linhas.append(f"| {nome[a['gravidade']]} | {_c(a['titulo'])} | {_c(a['detalhe'])} | {_c(a['onde'])} |")
    if not r["achados"]:
        linhas.append("| - | Nenhuma inconsistência encontrada pelas regras automáticas | - | - |")
    if externos:
        linhas += ["", "### 3. CONSULTAS DO FORNECEDOR", "| Base | Resultado |", "|---|---|"]
        linhas += [f"| {_c(e['base'])} | {_c(e['resultado'])} |" for e in externos]
    n = 4 if externos else 3
    linhas += ["", f"### {n}. PEÇAS IDENTIFICADAS", "| Peça | Onde | Data |", "|---|---|---|"]
    linhas += [f"| {ROTULO[d['tipo']]} | {_c(onde(p, d['indices']))} | {fmt(d['data'])} |" for d in r["docs"]]
    linhas += ["", f"### {n + 1}. LISTA DE VERIFICAÇÃO", "| Nº | Item | Situação | Onde |", "|---|---|---|---|"]
    linhas += [f"| {k} | {_c(c['texto'])} | {c['situacao']} | {_c(c['onde'])} |" for k, c in enumerate(r["checklist"], 1)]
    linhas += ["", f"### {n + 2}. LIMITES DA CONFERÊNCIA",
               "Conferência automática por regras objetivas sobre o texto dos arquivos. Não substitui a análise do "
               "agente de contratação nem o parecer jurídico. Páginas digitalizadas sem texto não são lidas; itens "
               "marcados como \"Verificar\" exigem conferência manual."]
    linhas += [f"- {a}" for a in r["avisos"]]
    return "\n".join(linhas)
