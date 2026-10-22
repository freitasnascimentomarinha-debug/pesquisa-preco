"""Serviços públicos compartilhados pelas páginas do AtaCotada (sem dependência do Streamlit).

Centraliza o que antes estava repetido nas páginas Consulta CNPJ, Detalhes Compra, Cotação Direta, Fornecedores e Notas Fiscais:
cliente HTTP com tentativas e erro padronizado (sempre com verificação de certificado), utilitários de CNPJ, busca no PNCP,
itens de ARP no Compras.gov.br, download dos CSVs de notas fiscais e o resumo correto do valor das notas.
"""

from __future__ import annotations

import datetime as _dt
import os
import re
import tempfile
import threading
import time
from typing import Any, Callable

import pandas as pd
import requests

USER_AGENT = {"User-Agent": "Mozilla/5.0"}
PNCP_SEARCH_URL = "https://pncp.gov.br/api/search/"
DADOSABERTOS_ARP_ITEM_URL = "https://dadosabertos.compras.gov.br/modulo-arp/2_consultarARPItem"
OPENCNPJ_URL = "https://api.opencnpj.org/{cnpj}"
BRASILAPI_CNPJ_URL = "https://brasilapi.com.br/api/cnpj/v1/{cnpj}"

# ── Cache em memória que só guarda sucessos (falha de API não fica "presa" por uma hora) ──────────────────────
_cache: dict[tuple, tuple[float, Any]] = {}
_cache_lock = threading.Lock()


def _com_cache(chave: tuple, ttl: int, buscar: Callable[[], tuple[Any, str | None]]) -> tuple[Any, str | None]:
    agora = time.time()
    with _cache_lock:
        achado = _cache.get(chave)
        if achado and agora - achado[0] < ttl:
            return achado[1], None
    dados, erro = buscar()
    if erro is None:
        with _cache_lock:
            _cache[chave] = (agora, dados)
    return dados, erro


# ── CNPJ ──────────────────────────────────────────────────────────────────────────────────────────────────────
def limpar_cnpj(valor: object) -> str:
    return re.sub(r"\D", "", str(valor or ""))


def formatar_cnpj(valor: object) -> str:
    """XX.XXX.XXX/XXXX-XX; se não tiver exatamente 14 dígitos, devolve o texto original."""
    d = limpar_cnpj(valor)
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}" if len(d) == 14 else str(valor or "")


def cnpj_valido(valor: object) -> bool:
    """14 dígitos, não todos iguais e com os dois dígitos verificadores corretos."""
    d = limpar_cnpj(valor)
    if len(d) != 14 or len(set(d)) == 1:
        return False
    for tamanho in (12, 13):
        pesos = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2] if tamanho == 12 else [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
        resto = sum(int(n) * p for n, p in zip(d[:tamanho], pesos)) % 11
        if int(d[tamanho]) != (0 if resto < 2 else 11 - resto):
            return False
    return True


# ── HTTP ──────────────────────────────────────────────────────────────────────────────────────────────────────
def http_get_json(url: str, params: dict | None = None, timeout: int = 30, tentativas: int = 2,
                  ok_404: bool = False) -> tuple[Any, str | None]:
    """GET que devolve (json, None) ou (None, mensagem de erro). 404 vira (None, None) quando `ok_404`.
    Sempre com verificação de certificado."""
    erro = "sem resposta"
    for tentativa in range(max(1, tentativas)):
        try:
            resposta = requests.get(url, params=params, timeout=timeout, headers=USER_AGENT)
            if resposta.status_code == 200:
                return resposta.json(), None
            if resposta.status_code == 404 and ok_404:
                return None, None
            erro = f"HTTP {resposta.status_code}"
            if resposta.status_code < 500 and resposta.status_code != 429:
                break
        except requests.exceptions.Timeout:
            erro = "tempo esgotado"
        except requests.exceptions.SSLError:
            erro = "falha na verificação do certificado"
            break
        except requests.exceptions.RequestException as exc:
            erro = f"falha de conexão ({type(exc).__name__})"
        except ValueError:
            erro = "resposta inválida (não é JSON)"
            break
        if tentativa + 1 < tentativas:
            time.sleep(0.6 * (tentativa + 1))
    return None, erro


# ── Cadastro do CNPJ ──────────────────────────────────────────────────────────────────────────────────────────
def consultar_cadastro_cnpj(cnpj: object) -> dict:
    """Dados cadastrais: OpenCNPJ (principal) e BrasilAPI (reserva de e-mail/telefone).
    Retorna {'opencnpj': dict|None, 'brasilapi': dict|None, 'erros': {fonte: mensagem}}."""
    digitos = limpar_cnpj(cnpj)
    resultado: dict = {"opencnpj": None, "brasilapi": None, "erros": {}}
    for chave, modelo in (("opencnpj", OPENCNPJ_URL), ("brasilapi", BRASILAPI_CNPJ_URL)):
        dados, erro = _com_cache((chave, digitos), 3600, lambda m=modelo: http_get_json(m.format(cnpj=digitos), timeout=10, tentativas=2, ok_404=True))
        resultado[chave] = dados if isinstance(dados, dict) else None
        if erro:
            resultado["erros"][chave] = erro
    return resultado


# ── PNCP (Search) ─────────────────────────────────────────────────────────────────────────────────────────────
def _data_do_registro(registro: dict) -> _dt.date | None:
    for campo in ("data_publicacao_pncp", "data_inicio_vigencia", "data_assinatura"):
        texto = str(registro.get(campo) or "")[:10]
        if texto:
            try:
                return _dt.date.fromisoformat(texto)
            except ValueError:
                continue
    return None


def pncp_search(cnpj: object, tipo_documento: str, data_min: _dt.date | None = None, max_paginas: int = 20) -> tuple[list[dict], str | None]:
    """Documentos do PNCP ligados a um CNPJ ('contrato', 'ata' ou 'compra'), do mais recente ao mais antigo.
    Com `data_min`, para de paginar ao passar da data e descarta o que for anterior (sem data, mantém).
    Retorna (itens, erro): em erro parcial, devolve o que conseguiu com a mensagem."""
    digitos = limpar_cnpj(cnpj)
    chave = ("pncp", digitos, tipo_documento, data_min, max_paginas)

    def buscar() -> tuple[list[dict], str | None]:
        itens: list[dict] = []
        vistos = 0
        for pagina in range(1, max_paginas + 1):
            dados, erro = http_get_json(PNCP_SEARCH_URL, {"q": digitos, "tipos_documento": tipo_documento, "pagina": pagina, "ordenacao": "-data"},
                                        timeout=60, tentativas=2)
            if erro:
                return itens, f"PNCP ({tipo_documento}): {erro}" + (f" — retornados {len(itens)} registro(s) antes da falha" if itens else "")
            pagina_itens = [i for i in (dados or {}).get("items", []) if isinstance(i, dict)]
            if not pagina_itens:
                break
            antigos = 0
            for item in pagina_itens:
                data = _data_do_registro(item) if data_min else None
                if data_min and data and data < data_min:
                    antigos += 1
                    continue
                itens.append(item)
            vistos += len(pagina_itens)
            if data_min and antigos == len(pagina_itens):
                break
            if vistos >= (dados or {}).get("total", 0):
                break
        return itens, None

    return _com_cache(chave, 3600, buscar)


# ── Compras.gov.br: itens de ARP por fornecedor ───────────────────────────────────────────────────────────────
def itens_arp_por_fornecedor(cnpj: object, anos_janela: int = 3, max_paginas: int = 20) -> tuple[list[dict], str | None]:
    """Itens de ARP em que o CNPJ é fornecedor registrado (2_consultarARPItem, niFornecedor).
    A API aceita janelas de até 365 dias, então varre em blocos de 360 dias. Retorna (itens, erro)."""
    digitos = limpar_cnpj(cnpj)

    def buscar() -> tuple[list[dict], str | None]:
        itens: list[dict] = []
        hoje = _dt.date.today()
        for i in range(anos_janela):
            d_fim = hoje - _dt.timedelta(days=360 * i)
            d_ini = hoje - _dt.timedelta(days=360 * (i + 1))
            bloco = 0
            for pagina in range(1, max_paginas + 1):
                dados, erro = http_get_json(DADOSABERTOS_ARP_ITEM_URL, {
                    "niFornecedor": digitos, "dataVigenciaInicialMin": d_ini.isoformat(), "dataVigenciaInicialMax": d_fim.isoformat(),
                    "pagina": pagina, "tamanhoPagina": 100}, timeout=60, tentativas=2)
                if erro:
                    return itens, f"Compras.gov.br (ARP): {erro}" + (f" — retornados {len(itens)} item(ns) antes da falha" if itens else "")
                resultado = (dados or {}).get("resultado", [])
                if not resultado:
                    break
                itens.extend(resultado)
                bloco += len(resultado)
                if bloco >= (dados or {}).get("totalRegistros", 0):
                    break
        return itens, None

    return _com_cache(("arp", digitos, anos_janela), 3600, buscar)


def arp_vigente_com_adesao(item: dict, hoje: _dt.date | None = None) -> bool:
    """Item de ARP ainda vigente e com quantidade máxima de adesão maior que zero."""
    limite = (hoje or _dt.date.today()).isoformat()
    return str(item.get("dataVigenciaFinal") or "")[:10] >= limite and (item.get("maximoAdesao") or 0) > 0


def separar_vigentes(itens: list[dict]) -> tuple[list[dict], list[dict]]:
    """(vigentes com adesão aberta, demais) — uma única passada."""
    vigentes: list[dict] = []
    outros: list[dict] = []
    hoje = _dt.date.today()
    for item in itens:
        (vigentes if arp_vigente_com_adesao(item, hoje) else outros).append(item)
    return vigentes, outros


# ── Notas fiscais (CSVs do Portal da Transparência no Google Drive) ───────────────────────────────────────────
NF_DOWNLOAD_URL = "https://drive.usercontent.google.com/download?id={file_id}&export=download&confirm=t"
NF_CACHE_DIR = os.path.join(tempfile.gettempdir(), "atacotada_nf")


def listar_arquivos_drive(folder_id: str) -> dict[str, str] | None:
    """{id: nome} dos arquivos da pasta pública do Drive, ou None se não conseguir listar."""
    try:
        resp = requests.get(f"https://drive.google.com/drive/folders/{folder_id}", timeout=30)
        resp.raise_for_status()
        ids = list(dict.fromkeys(re.findall(r'data-id="([a-zA-Z0-9_-]{20,})"', resp.text)))
    except requests.RequestException:
        return None
    arquivos: dict[str, str] = {}
    for fid in ids:
        nome = f"Arquivo {fid[:10]}"
        try:
            pagina = requests.get(f"https://drive.google.com/file/d/{fid}/view", timeout=15).text
            achado = re.search(r"<title>([^<]+)</title>", pagina)
            if achado:
                titulo = achado.group(1).replace(" - Google Drive", "").strip()
                if titulo and titulo != "Google Drive":
                    nome = titulo
        except requests.RequestException:
            pass
        arquivos[fid] = nome
    return arquivos or None


def baixar_csv_drive(file_id: str, progresso: Callable[[float, str], None] | None = None) -> str:
    """Baixa o CSV (cache local em arquivo) e devolve o caminho. Levanta RuntimeError se o Drive devolver uma página de erro."""
    os.makedirs(NF_CACHE_DIR, exist_ok=True)
    destino = os.path.join(NF_CACHE_DIR, f"{file_id}.csv")
    if os.path.exists(destino):
        if os.path.getsize(destino) > 10000:
            return destino
        os.remove(destino)

    resposta = requests.get(NF_DOWNLOAD_URL.format(file_id=file_id), stream=True, timeout=600)
    resposta.raise_for_status()
    if "text/html" in resposta.headers.get("Content-Type", ""):
        resposta.close()
        raise RuntimeError("O Google Drive retornou uma página de erro (quota de downloads excedida).")

    total = int(resposta.headers.get("content-length", 0))
    baixado = 0
    temporario = destino + ".tmp"
    with open(temporario, "wb") as arquivo:
        for bloco in resposta.iter_content(chunk_size=131072):
            arquivo.write(bloco)
            baixado += len(bloco)
            if progresso and total:
                fracao = min(baixado / total, 1.0)
                progresso(fracao, f"Baixando... {baixado / 1048576:.0f} / {total / 1048576:.0f} MB ({fracao * 100:.0f}%)")
    with open(temporario, "rb") as arquivo:
        inicio = arquivo.read(200)
    if b"<!DOCTYPE" in inicio or b"<html" in inicio.lower():
        os.remove(temporario)
        raise RuntimeError("O arquivo baixado não é um CSV válido (quota do Google Drive excedida).")
    os.replace(temporario, destino)
    return destino


def pesquisar_notas_por_cnpj(caminho: str, cnpj: object, max_resultados: int = 1000) -> tuple[pd.DataFrame, int]:
    """Linhas do CSV em que o emitente é o CNPJ. Retorna (DataFrame, total de linhas lidas)."""
    alvo = limpar_cnpj(cnpj)
    achados: list[pd.DataFrame] = []
    lidas = 0
    for bloco in pd.read_csv(caminho, sep=";", encoding="latin-1", chunksize=50000, dtype=str, on_bad_lines="skip"):
        if "CPF/CNPJ Emitente" in bloco.columns:
            limpo = bloco["CPF/CNPJ Emitente"].astype(str).str.replace(r"\D", "", regex=True)
            filtrado = bloco[limpo == alvo]
            if len(filtrado):
                achados.append(filtrado)
        lidas += len(bloco)
        if sum(len(a) for a in achados) >= max_resultados:
            break
    if achados:
        return pd.concat(achados, ignore_index=True).head(max_resultados), lidas
    return pd.DataFrame(), lidas


def numero_br(valor: object) -> float | None:
    """'1.234,56' -> 1234.56; '1234.56' -> 1234.56; vazio/inválido -> None."""
    texto = str(valor if valor is not None else "").strip()
    if not texto or texto.lower() == "nan":
        return None
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def resumo_notas(df: pd.DataFrame) -> dict:
    """Totais corretos de um resultado de notas fiscais por item.

    O CSV traz uma linha por ITEM; o valor total da nota se repete em todas as linhas da mesma nota. Por isso:
    - notas = chaves de acesso distintas (ou, sem a coluna, número+série+período);
    - valor = soma dos valores dos itens (coluna de item) ou, só se não houver, valor da nota contado uma vez por nota.
    Retorna {'itens', 'notas', 'valor', 'base'} (valor None se nenhuma coluna de valor foi reconhecida)."""
    if df is None or df.empty:
        return {"itens": 0, "notas": 0, "valor": 0.0, "base": "itens"}

    col_chave = next((c for c in ("CHAVE DE ACESSO", "CHAVE DE ACESSO NFE") if c in df.columns), None)
    if col_chave:
        chave = df[col_chave].astype(str)
    else:
        partes = [df[c].astype(str) for c in ("NÚMERO", "SÉRIE", "_arquivo_origem") if c in df.columns]
        chave = partes[0].str.cat(partes[1:], sep="|") if partes else pd.Series(range(len(df)), index=df.index).astype(str)
    notas = int(chave.nunique())

    col_item = next((c for c in ("VALOR TOTAL", "VALOR TOTAL DO ITEM", "VALOR TOTAL DO PRODUTO") if c in df.columns), None)
    if col_item:
        valores = df[col_item].map(numero_br)
        return {"itens": len(df), "notas": notas, "valor": float(valores.fillna(0).sum()), "base": "itens"}

    col_nota = next((c for c in ("VALOR TOTAL DA NOTA", "VALOR NOTA FISCAL") if c in df.columns), None)
    if col_nota:
        unico = df.assign(_chave=chave.values).drop_duplicates("_chave")
        return {"itens": len(df), "notas": notas, "valor": float(unico[col_nota].map(numero_br).fillna(0).sum()), "base": "notas"}

    return {"itens": len(df), "notas": notas, "valor": None, "base": "itens"}


# ── Exportação ────────────────────────────────────────────────────────────────────────────────────────────────
def gerar_excel(abas: dict[str, pd.DataFrame]) -> bytes:
    """Excel (.xlsx) com uma aba por DataFrame não vazio: cabeçalho em negrito, filtro, linha fixa e colunas ajustadas."""
    from io import BytesIO

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    saida = BytesIO()
    with pd.ExcelWriter(saida, engine="openpyxl") as escritor:
        escritas = 0
        for nome, df in abas.items():
            if df is None or df.empty:
                continue
            aba = re.sub(r"[\[\]\*\?/\\:]", "-", nome)[:31]
            df.to_excel(escritor, sheet_name=aba, index=False)
            ws = escritor.sheets[aba]
            for celula in ws[1]:
                celula.font = Font(bold=True, color="FFFFFF")
                celula.fill = PatternFill("solid", fgColor="1F3864")
                celula.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, coluna in enumerate(df.columns, start=1):
                maior = max([len(str(coluna))] + [len(str(v)) for v in df[coluna].head(300)])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, maior + 2), 60)
            escritas += 1
        if not escritas:
            pd.DataFrame({"Aviso": ["Nenhum registro para exportar."]}).to_excel(escritor, sheet_name="Resultado", index=False)
    return saida.getvalue()


# ── Sanções (CEIS, CNEP, CEPIM): bases diárias públicas do Portal da Transparência (sem chave de API) ─────────
SANCOES_URL = "https://portaldatransparencia.gov.br/download-de-dados/{tipo}/{data}"
SANCOES_CACHE_DIR = os.path.join(tempfile.gettempdir(), "atacotada_sancoes")
SANCOES_TIPOS = {
    "ceis": ("CEIS", "Cadastro de Empresas Inidôneas e Suspensas", "CPF OU CNPJ DO SANCIONADO"),
    "cnep": ("CNEP", "Cadastro Nacional de Empresas Punidas (Lei Anticorrupção)", "CPF OU CNPJ DO SANCIONADO"),
    "cepim": ("CEPIM", "Cadastro de Entidades Privadas sem Fins Lucrativos Impedidas", "CNPJ ENTIDADE"),
}


_ultima_sondagem: dict[str, float] = {}


def _baixar_base_sancoes(tipo: str, dias_atras: int = 12) -> tuple[str | None, str | None]:
    """Caminho do CSV mais recente da base (com cache em disco por data) e a data dela (AAAA-MM-DD)."""
    import zipfile

    os.makedirs(SANCOES_CACHE_DIR, exist_ok=True)
    hoje = _dt.date.today()
    # a base é atualizada em dias úteis: uma cópia local de até 2 dias dispensa nova ida ao portal
    locais = sorted(n for n in os.listdir(SANCOES_CACHE_DIR) if n.startswith(f"{tipo}_") and n.endswith(".csv"))
    if locais:
        dia_local = locais[-1][len(tipo) + 1:-4]
        try:
            recente = (hoje - _dt.datetime.strptime(dia_local, "%Y%m%d").date()).days <= 2
        except ValueError:
            recente = False
        if recente or time.time() - _ultima_sondagem.get(tipo, 0) < 6 * 3600:  # base mais nova não existe: não insiste a cada consulta
            return os.path.join(SANCOES_CACHE_DIR, locais[-1]), dia_local
    _ultima_sondagem[tipo] = time.time()
    for recuo in range(dias_atras + 1):
        dia = (hoje - _dt.timedelta(days=recuo)).strftime("%Y%m%d")
        destino = os.path.join(SANCOES_CACHE_DIR, f"{tipo}_{dia}.csv")
        if os.path.exists(destino):
            return destino, dia
        try:
            resposta = requests.get(SANCOES_URL.format(tipo=tipo, data=dia), headers=USER_AGENT, timeout=90)
        except requests.RequestException:
            continue
        if resposta.status_code != 200 or "zip" not in resposta.headers.get("content-type", ""):
            continue
        try:
            from io import BytesIO

            with zipfile.ZipFile(BytesIO(resposta.content)) as arquivo:
                nome = next(n for n in arquivo.namelist() if n.lower().endswith(".csv"))
                dados = arquivo.read(nome)
        except (zipfile.BadZipFile, StopIteration):
            continue
        with open(destino + ".tmp", "wb") as saida:
            saida.write(dados)
        os.replace(destino + ".tmp", destino)
        for antigo in os.listdir(SANCOES_CACHE_DIR):  # mantém só a base mais recente de cada tipo
            if antigo.startswith(f"{tipo}_") and antigo != os.path.basename(destino):
                try:
                    os.remove(os.path.join(SANCOES_CACHE_DIR, antigo))
                except OSError:
                    pass
        return destino, dia
    return None, None


_bases_lidas: dict[str, pd.DataFrame] = {}  # CSV já lido em memória (uma base por tipo)


def consultar_sancoes(cnpj: object) -> tuple[dict[str, list[dict]], dict[str, str], str | None]:
    """Sanções do CNPJ no CEIS, CNEP e CEPIM.
    Retorna ({'CEIS': [registro, ...], ...}, {'CEIS': 'AAAA-MM-DD' (data da base), ...}, erro).
    Cada registro é o dicionário das colunas do CSV, com 'vigente' (bool) calculado pelo fim da sanção."""
    alvo = limpar_cnpj(cnpj)
    resultado: dict[str, list[dict]] = {}
    datas: dict[str, str] = {}
    erros: list[str] = []
    hoje = _dt.date.today()
    for tipo, (sigla, _descricao, coluna) in SANCOES_TIPOS.items():
        def buscar(tipo=tipo, coluna=coluna) -> tuple[Any, str | None]:
            caminho, dia = _baixar_base_sancoes(tipo)
            if not caminho:
                return None, f"base {tipo.upper()} indisponível no Portal da Transparência"
            df = _bases_lidas.get(caminho)
            if df is None:
                try:
                    df = pd.read_csv(caminho, sep=";", encoding="latin-1", dtype=str, keep_default_na=False)
                except (OSError, ValueError, pd.errors.ParserError) as exc:
                    return None, f"base {tipo.upper()} ilegível ({type(exc).__name__})"
                for velho in [c for c in _bases_lidas if os.path.basename(c).startswith(f"{tipo}_")]:
                    del _bases_lidas[velho]
                _bases_lidas[caminho] = df
            achados = df[df[coluna].str.replace(r"\D", "", regex=True) == alvo]
            return (achados.to_dict("records"), dia), None

        dados, erro = _com_cache(("sancoes", tipo, alvo, hoje), 6 * 3600, buscar)
        if erro:
            erros.append(erro)
            continue
        registros, dia = dados
        for registro in registros:
            fim = str(registro.get("DATA FINAL SANÇÃO", "")).strip()
            try:
                registro["vigente"] = (not fim) or _dt.datetime.strptime(fim, "%d/%m/%Y").date() >= hoje
            except ValueError:
                registro["vigente"] = True  # data ilegível: trata como vigente por prudência
        resultado[sigla] = registros
        datas[sigla] = f"{dia[:4]}-{dia[4:6]}-{dia[6:]}"
    return resultado, datas, "; ".join(erros) or None


# ── Compras.gov.br: UASG e cadastro de fornecedores (SICAF) ───────────────────────────────────────────────────
DADOSABERTOS = "https://dadosabertos.compras.gov.br"


def dados_uasg(codigo_uasg: object) -> tuple[dict | None, str | None]:
    """Cadastro da UASG (órgão, CNPJ, UF, município, unidade espelho). Tenta ativa e depois inativa. (None, None) se não existir."""
    codigo = re.sub(r"\D", "", str(codigo_uasg or ""))
    if not codigo:
        return None, None

    def buscar() -> tuple[Any, str | None]:
        ultimo_erro = None
        for status in ("true", "false"):
            dados, erro = http_get_json(f"{DADOSABERTOS}/modulo-uasg/1_consultarUasg",
                                        {"codigoUasg": codigo, "statusUasg": status, "pagina": 1}, timeout=40, tentativas=3)
            if erro:
                ultimo_erro = erro
                continue
            resultado = (dados or {}).get("resultado") or []
            if resultado:
                return resultado[0], None
        return None, ultimo_erro

    return _com_cache(("uasg", codigo), 24 * 3600, buscar)


def consultar_sicaf(cnpj: object) -> tuple[dict, str | None]:
    """Situação do fornecedor no cadastro do Compras.gov.br (base do SICAF).
    Retorna ({'encontrado', 'ativo', 'habilitado', 'razao_social', 'cnae', 'porte', 'natureza', 'municipio', 'uf'}, erro).
    Mostra o cadastro (ativo e habilitado a licitar), não a validade das certidões, que exige acesso ao SICAF."""
    digitos = limpar_cnpj(cnpj)
    vazio = {"encontrado": False, "ativo": None, "habilitado": None}
    if len(digitos) not in (11, 14):
        return vazio, None
    campo = "cnpj" if len(digitos) == 14 else "cpf"

    def buscar() -> tuple[Any, str | None]:
        for ativo in ("true", "false"):
            dados, erro = http_get_json(f"{DADOSABERTOS}/modulo-fornecedor/1_consultarFornecedor",
                                        {campo: digitos, "ativo": ativo, "pagina": 1, "tamanhoPagina": 10}, timeout=40, tentativas=3)
            if erro:
                return None, f"Compras.gov.br (fornecedor): {erro}"
            resultado = (dados or {}).get("resultado") or []
            if resultado:
                r = resultado[0]
                return {"encontrado": True, "ativo": bool(r.get("ativo")), "habilitado": bool(r.get("habilitadoLicitar")),
                        "razao_social": r.get("nomeRazaoSocialFornecedor") or "",
                        "cnae": f"{r.get('codigoCnae') or ''} - {r.get('nomeCnae') or ''}".strip(" -"),
                        "porte": r.get("porteEmpresaNome") or "", "natureza": r.get("naturezaJuridicaNome") or "",
                        "municipio": r.get("nomeMunicipio") or "", "uf": r.get("ufSigla") or ""}, None
        return dict(vazio), None

    dados, erro = _com_cache(("sicaf", digitos), 6 * 3600, buscar)
    return (dados or dict(vazio)), erro


def texto_sicaf(situacao: dict) -> str:
    """Rótulo curto da situação no cadastro de fornecedores."""
    if not situacao or situacao.get("encontrado") is False:
        return "Não cadastrado"
    if situacao.get("ativo") and situacao.get("habilitado"):
        return "Ativo e habilitado a licitar"
    if situacao.get("ativo"):
        return "Ativo, não habilitado a licitar"
    return "Cadastro inativo"
