"""Relatórios (PDF e Excel) da pesquisa em lote de notas fiscais e dos fornecedores encontrados.

Seguem o layout da Cotação Rápida: mapa comparativo de preços na primeira página, uma página por item e a justificativa.
"""

from __future__ import annotations

import io

import pandas as pd
from fpdf import FPDF
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import relatorio_cotacao_rapida as base
from cotacao_rapida import MAX_PRECOS, MIN_PRECOS, TOLERANCIA
from fornecedores_nf import formatar_cnpj
from relatorio_cotacao_rapida import (
    AZUL, AZUL_CARTAO, DOURADO, JUSTIFICATIVA_COTACAO, LARGURA, PDFCotacaoRapida, STATUS_TEXTO, _cabecalho_tabela, _cartoes, _quebrar,
    _seguro, _truncar, moeda,
)

STATUS_TEXTO["sem_notas"] = "Sem notas fiscais encontradas"
TITULO_PDF = "Pesquisa de Notas Fiscais - Mapa Comparativo de Precos"
TITULO_FORNECEDORES = "Fornecedores encontrados - Notas Fiscais"


def metodologia(info: dict) -> str:
    periodo = ", ".join(info.get("arquivos", [])) or "arquivos selecionados"
    return (
        "Metodologia da Pesquisa em Lote de Notas Fiscais\n\n"
        "1. Fonte: dados abertos de notas fiscais eletrônicas (NF-e, itens de nota) do Portal da Transparência, "
        f"arquivos consultados: {periodo}. Foram analisadas {info.get('linhas', 0):,} linhas de itens de nota.".replace(",", ".") + "\n\n"
        "2. Para cada descrição informada foram localizadas as notas cujo item contém todas as palavras da descrição e as "
        "especificações pedidas (ex.: A4 equivale a 210 x 297 mm), sem diferenciar acentos, maiúsculas ou plural. "
        f"Filtros aplicados: {info.get('filtros', 'nenhum')}.\n\n"
        "3. Os preços unitários foram comparados apenas dentro da mesma unidade (unidades equivalentes, como UN/UND/UNID, "
        "foram agrupadas) e, quando indicado, foram descartados preços inexequíveis ou extremos (abaixo de 30% ou acima de 300% "
        "da mediana) e os valores atípicos (outliers) pelo método do intervalo interquartil (IQR).\n\n"
        f"4. Foram selecionados até {MAX_PRECOS} preços por item, priorizando fornecedores distintos e as notas mais recentes, de modo que "
        f"cada preço listado não se afaste mais que {TOLERANCIA * 100:.0f}% (para mais ou para menos) da média dos preços listados. "
        f"Itens com menos de {MIN_PRECOS} preços são sinalizados e devem ser complementados por outras fontes (IN SEGES/ME nº 65/2021).\n\n"
        "5. A descrição do item na nota fiscal é texto livre do emitente: marca, quantidade por embalagem e demais características "
        "podem variar. O requisitante deve conferir se o produto da nota corresponde ao objeto pretendido antes de utilizar o valor. "
        "A chave de acesso (44 dígitos) de cada nota consta da planilha Excel para consulta do documento original no Portal da NF-e."
    )


def _pagina_item_nf(info: dict):
    def pagina(pdf: FPDF, numero: int, r: dict) -> None:
        pdf.add_page()
        pdf.set_y(35)
        pdf.set_fill_color(*AZUL_CARTAO)
        pdf.set_text_color(*DOURADO)
        pdf.set_font("Helvetica", "B", 8)
        pdf.cell(LARGURA, 6, f"  ITEM {numero} - {STATUS_TEXTO.get(r['status'], '').upper()}", 0, 1, "L", True)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 9)
        for linha in _quebrar(pdf, r["descricao"], LARGURA - 4, 3):
            pdf.cell(LARGURA, 5, "  " + linha, 0, 1, "L", True)
        pdf.set_font("Helvetica", "", 6.5)
        pdf.set_text_color(70, 70, 70)
        pdf.cell(0, 5, _seguro(f"{r['notas']} notas fiscais, {r['fornecedores']} fornecedores. Filtros: {info.get('filtros', 'nenhum')}"), ln=True)
        pdf.ln(1)
        if not r["precos"]:
            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(150, 30, 30)
            texto = ("Nenhuma nota fiscal encontrada com os filtros informados." if r["status"] == "sem_notas"
                     else "Os precos encontrados nao formaram um conjunto coerente (a ate 30% da media) apos os filtros.")
            pdf.multi_cell(LARGURA, 5, _seguro(texto), new_x="LMARGIN", new_y="NEXT")
            return
        s = r["stats"]
        _cartoes(pdf, [("MINIMO", f"R$ {moeda(s['min'])}"), ("MEDIA UNITARIA", f"R$ {moeda(s['media'])}"), ("MEDIANA", f"R$ {moeda(s['mediana'])}"),
                       ("MAXIMO", f"R$ {moeda(s['max'])}"), ("DESVIO PADRAO", f"R$ {moeda(s['desvio'])}"), ("CV", f"{s['cv']:.2f}%".replace(".", ","))], pdf.get_y())
        colunas = [("Data", 16), ("Fornecedor (emitente)", 56), ("CNPJ", 28), ("UF", 7), ("Municipio", 28), ("Unid.", 12), ("Qtd", 14),
                   ("V. Unitario", 20), ("Destinatario", 50), ("UF", 7), ("Produto na nota", 43)]
        _cabecalho_tabela(pdf, colunas)
        pdf.set_font("Helvetica", "", 6)
        pdf.set_text_color(51, 51, 51)
        for indice, p in enumerate(r["precos"]):
            valores = [str(p["data"] or "")[:10], p["fornecedor"], formatar_cnpj(p["cnpj"]), p["uf"], p["municipio"], p["sigla"],
                       f"{p['quantidade']:g}" if p["quantidade"] is not None else "", f"R$ {moeda(p['preco'])}", p["nome_uasg"], p["uf_dest"], p["descricao"]]
            pdf.set_fill_color(245, 248, 255)
            for (_, largura), valor in zip(colunas, valores):
                pdf.cell(largura, 6, _truncar(pdf, valor, largura - 2), 1, 0, "C", indice % 2 == 1)
            pdf.ln()
        pdf.ln(3)
        pdf.set_font("Helvetica", "", 7)
        pdf.set_text_color(70, 70, 70)
        validos = r["faixa_validos"]
        notas = [
            f"Unidade considerada: {r['unidade']}.",
            f"Universo analisado: {r['universo']} itens de nota; outliers/precos inexequiveis removidos: {r['outliers']}"
            + (f" (faixa valida: R$ {moeda(validos[0])} a R$ {moeda(validos[1])})." if validos else "."),
            f"Os {len(r['precos'])} precos listados ficam a ate {TOLERANCIA * 100:.0f}% da media (faixa de R$ {moeda(s['media'] * (1 - TOLERANCIA))} a R$ {moeda(s['media'] * (1 + TOLERANCIA))}).",
        ]
        if r["status"] == "insuficiente":
            notas.append(f"ATENCAO: menos de {MIN_PRECOS} precos - complementar com outras fontes (IN 65/2021).")
        for nota in notas:
            pdf.multi_cell(LARGURA, 4, _seguro(nota), new_x="LMARGIN", new_y="NEXT")

    return pagina


def gerar_pdf_mapa(resultados: list[dict], info: dict) -> bytes:
    return base.gerar_pdf(
        resultados,
        titulo=TITULO_PDF,
        pagina_item=_pagina_item_nf(info),
        textos=(JUSTIFICATIVA_COTACAO, metodologia(info)),
        argumentos_mapa={
            "rotulo_col3": "Notas / fornecedores",
            "linhas_col3": lambda r: [f"{r['notas']} notas", f"{r['fornecedores']} fornecedores"] if r["notas"] else ["-"],
            "fonte": "Fonte: Portal da Transparencia - notas fiscais eletronicas (itens de nota), mesma unidade, outliers e precos inexequiveis removidos. "
                     "Detalhamento de cada item e chaves de acesso (Excel) nas paginas seguintes.",
        },
    )


def tabela_mapa_nf(resultados: list[dict]) -> pd.DataFrame:
    linhas = []
    for numero, r in enumerate(resultados, start=1):
        precos = [p["preco"] for p in r["precos"]]
        s = r["stats"]
        linha = {"Item": numero, "Descrição pesquisada": r["descricao"], "Notas encontradas": r["notas"], "Fornecedores": r["fornecedores"], "Unidade": r["unidade"] or "-"}
        for n in range(MAX_PRECOS):
            linha[f"Preço {n + 1}"] = precos[n] if n < len(precos) else None
        linha.update({
            "Média unitária": s["media"] if s else None, "Mediana": s["mediana"] if s else None, "Mínimo": s["min"] if s else None,
            "Máximo": s["max"] if s else None, "Desvio padrão": s["desvio"] if s else None, "CV (%)": s["cv"] if s else None,
            "Nº de preços": len(precos), "Situação": STATUS_TEXTO.get(r["status"], ""),
        })
        linhas.append(linha)
    return pd.DataFrame(linhas)


def _formatar_planilhas(escritor: pd.ExcelWriter, colunas_moeda: tuple[str, ...]) -> None:
    for planilha in escritor.book.worksheets:
        for celula in planilha[1]:
            celula.font = Font(color="FFFFFF", bold=True)
            celula.fill = PatternFill("solid", fgColor="001A4D")
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for coluna in planilha.columns:
            maior = max((len(str(c.value or "")) for c in coluna[:300]), default=10)
            planilha.column_dimensions[get_column_letter(coluna[0].column)].width = min(max(maior + 2, 11), 70)
            titulo = str(coluna[0].value or "")
            if titulo.startswith("Preço") or titulo in colunas_moeda:
                for celula in coluna[1:]:
                    celula.number_format = '"R$" #,##0.00'
        planilha.freeze_panes = "A2"


def gerar_excel_mapa(resultados: list[dict], info: dict) -> bytes:
    selecionados, todas = [], []
    for numero, r in enumerate(resultados, start=1):
        usados = {(p["id_compra"], p["id_item"]) for p in r["precos"]}
        for p in r["precos"]:
            selecionados.append({"Item": numero, "Descrição pesquisada": r["descricao"], "Data": str(p["data"])[:10], "Fornecedor": p["fornecedor"],
                                 "CNPJ": formatar_cnpj(p["cnpj"]), "UF emitente": p["uf"], "Município": p["municipio"], "Unidade": p["sigla"], "Quantidade": p["quantidade"],
                                 "Valor unitário": p["preco"], "Destinatário": p["nome_uasg"], "UF destinatário": p["uf_dest"], "Produto na nota": p["descricao"],
                                 "NCM/SH": p["ncm"], "Natureza da operação": p["natureza"], "Chave de acesso": p["id_compra"]})
        for p in r.get("registros", [])[:5000]:
            todas.append({"Item": numero, "Descrição pesquisada": r["descricao"], "Usado no mapa": "Sim" if (p["id_compra"], p["id_item"]) in usados else "Não",
                          "Data": str(p["data"])[:10], "Fornecedor": p["fornecedor"], "CNPJ": formatar_cnpj(p["cnpj"]), "UF emitente": p["uf"], "Unidade": p["sigla"],
                          "Quantidade": p["quantidade"], "Valor unitário": p["preco"], "Destinatário": p["nome_uasg"], "UF destinatário": p["uf_dest"],
                          "Produto na nota": p["descricao"], "Chave de acesso": p["id_compra"]})
    saida = io.BytesIO()
    with pd.ExcelWriter(saida, engine="openpyxl") as escritor:
        tabela_mapa_nf(resultados).to_excel(escritor, index=False, sheet_name="Mapa Comparativo")
        pd.DataFrame(selecionados).to_excel(escritor, index=False, sheet_name="Preços selecionados")
        pd.DataFrame(todas).to_excel(escritor, index=False, sheet_name="Todas as notas")
        pd.DataFrame({"Parâmetro": ["Arquivos consultados", "Linhas analisadas", "Filtros", "Preços por item", "Tolerância sobre a média", "Mínimo recomendado de preços"],
                      "Valor": [", ".join(info.get("arquivos", [])), info.get("linhas", 0), info.get("filtros", "nenhum"), MAX_PRECOS, f"{TOLERANCIA * 100:.0f}%", MIN_PRECOS]}
                     ).to_excel(escritor, index=False, sheet_name="Parâmetros")
        _formatar_planilhas(escritor, ("Média unitária", "Mediana", "Mínimo", "Máximo", "Desvio padrão", "Valor unitário"))
    return saida.getvalue()


# ------------------------------------------------------------------ fornecedores

COLUNAS_PDF_FORNECEDORES = [("CNPJ", 27, "CNPJ"), ("Razao social", 48, "Razão social"), ("UF", 8, "UF"), ("Telefones", 32, "Telefones"), ("E-mail", 42, "E-mail"),
                            ("CNAE principal", 45, "CNAE principal"), ("Natureza dos itens que vende", 79, "Natureza dos itens que vende")]


def gerar_pdf_fornecedores(tabela: pd.DataFrame, info: dict) -> bytes:
    pdf = PDFCotacaoRapida(orientation="L", unit="mm", format="A4", titulo=TITULO_FORNECEDORES)
    pdf.set_margins(8, 35, 8)
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=False)
    pdf.add_page()
    _cartoes(pdf, [("FORNECEDORES", str(len(tabela))), ("COM E-MAIL", str(int((tabela["E-mail"] != "Não informado").sum()))),
                   ("COM TELEFONE", str(int((tabela["Telefones"] != "Não informado").sum()))),
                   ("ESTADOS", str(tabela["UF"].replace("Não informado", pd.NA).nunique()))], 38)
    pdf.set_font("Helvetica", "", 6.5)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 4, _seguro(f"Fornecedores que emitiram as notas encontradas ({info.get('filtros', 'sem filtros')}). Dados cadastrais: OpenCNPJ (Receita Federal)."), ln=True)
    pdf.ln(1)
    colunas = [(rotulo, largura) for rotulo, largura, _ in COLUNAS_PDF_FORNECEDORES]
    _cabecalho_tabela(pdf, colunas)
    for numero, (_, linha) in enumerate(tabela.iterrows()):
        pdf.set_font("Helvetica", "", 6.3)
        celulas = [_quebrar(pdf, str(linha[campo]), largura - 2, 6 if campo.startswith("Natureza") else 3) for _, largura, campo in COLUNAS_PDF_FORNECEDORES]
        altura = max(7.0, 3.4 * max(len(c) for c in celulas) + 2)
        if pdf.get_y() + altura > pdf.h - 14:
            pdf.add_page()
            pdf.set_y(35)
            _cabecalho_tabela(pdf, colunas)
        x, y = pdf.l_margin, pdf.get_y()
        for (_, largura, _), linhas in zip(COLUNAS_PDF_FORNECEDORES, celulas):
            pdf.set_fill_color(*((245, 248, 255) if numero % 2 else (255, 255, 255)))
            pdf.set_draw_color(190, 190, 190)
            pdf.rect(x, y, largura, altura, "DF")
            pdf.set_text_color(51, 51, 51)
            for n, texto in enumerate(linhas):
                pdf.set_xy(x + 1, y + 1 + 3.4 * n)
                pdf.cell(largura - 2, 3.4, texto, 0, 0, "L")
            x += largura
        pdf.set_y(y + altura)
    return bytes(pdf.output())


def gerar_excel_fornecedores(tabela: pd.DataFrame) -> bytes:
    saida = io.BytesIO()
    with pd.ExcelWriter(saida, engine="openpyxl") as escritor:
        tabela.to_excel(escritor, index=False, sheet_name="Fornecedores")
        _formatar_planilhas(escritor, ())
        planilha = escritor.book["Fornecedores"]
        planilha.auto_filter.ref = planilha.dimensions
        for linha in planilha.iter_rows(min_row=2):
            for celula in linha:
                celula.alignment = Alignment(vertical="top", wrap_text=True)
    return saida.getvalue()
