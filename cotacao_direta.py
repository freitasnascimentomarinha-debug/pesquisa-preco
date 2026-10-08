"""Cotação Direta: pedido de cotação por e-mail aos fornecedores que venderam recentemente à Administração.

Fluxo: lista de itens -> CATMAT sugerido (automático) -> fornecedores recentes por item (Compras.gov, sem repetir fornecedor entre itens)
-> proposta em Word + e-mail padrão -> envio único com os fornecedores em cópia oculta (Resend).
Este módulo só tem contas e chamadas de rede; a tela fica em pages/Cotação_Direta.py.
"""

from __future__ import annotations

import base64
import datetime as dt
import hmac
import html
import io
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests

from lista_itens import formatar_quantidade, interpretar_linha, interpretar_lista, itens_do_dataframe  # noqa: F401  (a lista de itens é compartilhada com as outras páginas)

# ---------- parâmetros ----------
FORNECEDORES_POR_ITEM = 2
MIN_FORNECEDORES = 3  # IN SEGES/ME nº 65/2021: ao menos três fornecedores consultados (todos recebem todos os itens)
LIMIAR_CATMAT = 70.0  # % mínimo de combinação para sugerir um CATMAT
MAX_CATMAT_SUGERIDOS = 3
JANELA_DIAS = 365
MAX_CONSULTAS_CONTATO_POR_ITEM = 40  # CNPJs consultados (OpenCNPJ) por item até achar fornecedores com e-mail
LOTE_BCC = 45  # destinatários ocultos por mensagem (o Resend aceita até 50 por envio, contando o destinatário visível)
PRAZO_PADRAO_DIAS_UTEIS = 5
VALIDADE_PADRAO_DIAS = 60
API_RESEND = "https://api.resend.com/emails"
SITUACOES_INATIVAS = ("baix", "inapt", "suspens", "nul")
EMAIL_VALIDO = re.compile(r"^[A-Za-z0-9._%+\-']+@[A-Za-z0-9\-]+(\.[A-Za-z0-9\-]+)+$")


# ---------- OM e CEP ----------

def somente_digitos(texto: object) -> str:
    return re.sub(r"\D", "", str(texto or ""))


def formatar_cep(cep: object) -> str:
    digitos = somente_digitos(cep)
    return f"{digitos[:5]}-{digitos[5:8]}" if len(digitos) == 8 else str(cep or "")


def consultar_cep(cep: str) -> dict[str, str] | None:
    """Endereço pelo CEP (ViaCEP; BrasilAPI de reserva). None se não achar."""
    digitos = somente_digitos(cep)
    if len(digitos) != 8:
        return None
    try:
        resposta = requests.get(f"https://viacep.com.br/ws/{digitos}/json/", timeout=8)
        if resposta.status_code == 200:
            dados = resposta.json()
            if not dados.get("erro"):
                return {"logradouro": dados.get("logradouro", ""), "bairro": dados.get("bairro", ""),
                        "cidade": dados.get("localidade", ""), "uf": dados.get("uf", "")}
    except (requests.RequestException, ValueError):
        pass
    try:
        resposta = requests.get(f"https://brasilapi.com.br/api/cep/v1/{digitos}", timeout=8)
        if resposta.status_code == 200:
            dados = resposta.json()
            return {"logradouro": dados.get("street", ""), "bairro": dados.get("neighborhood", ""),
                    "cidade": dados.get("city", ""), "uf": dados.get("state", "")}
    except (requests.RequestException, ValueError):
        pass
    return None


def formatar_cnpj(cnpj: object) -> str:
    d = somente_digitos(cnpj)
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:14]}" if len(d) == 14 else str(cnpj or "")


def cnpj_valido(cnpj: object) -> bool:
    """14 dígitos, não todos iguais e com os dois dígitos verificadores corretos."""
    d = somente_digitos(cnpj)
    if len(d) != 14 or len(set(d)) == 1:
        return False
    for tamanho in (12, 13):
        pesos = ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2] if tamanho == 12 else [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
        resto = sum(int(n) * p for n, p in zip(d[:tamanho], pesos)) % 11
        if int(d[tamanho]) != (0 if resto < 2 else 11 - resto):
            return False
    return True


def endereco_completo(om: dict[str, str]) -> str:
    partes = [om.get("logradouro", "").strip(), om.get("numero", "").strip(), om.get("complemento", "").strip()]
    rua = ", ".join(p for p in partes if p)
    cidade = "/".join(p for p in (om.get("cidade", "").strip(), om.get("uf", "").strip()) if p)
    return " – ".join(p for p in (rua, om.get("bairro", "").strip(), cidade, f"CEP {formatar_cep(om.get('cep'))}" if om.get("cep") else "") if p)


CAMPOS_OM_SALVOS = ("nome", "cnpj", "cep", "logradouro", "numero", "complemento", "bairro", "cidade", "uf")  # telefone, e-mail e responsável nunca são guardados


def chave_om(nome: str) -> str:
    """Nome da OM sem acento, maiúsculas e espaços repetidos: 'Centro  de Operações' = 'CENTRO DE OPERACOES'."""
    import unicodedata

    sem_acento = "".join(c for c in unicodedata.normalize("NFD", str(nome or "")) if unicodedata.category(c) != "Mn")
    return " ".join(sem_acento.upper().split())


def registrar_om(memoria: dict, om: dict[str, str]) -> bool:
    """Guarda (ou atualiza) o nome e o endereço da OM na memória. Devolve True se algo mudou. Telefone, e-mail e responsável ficam de fora."""
    chave = chave_om(om.get("nome", ""))
    if not chave:
        return False
    novo = {campo: str(om.get(campo, "") or "").strip() for campo in CAMPOS_OM_SALVOS}
    oms = memoria.setdefault("oms", {})
    mudou = {k: v for k, v in oms.get(chave, {}).items() if k != "ultimo_uso"} != novo
    oms[chave] = {**novo, "ultimo_uso": dt.date.today().isoformat()}
    if mudou:
        memoria["_mudou"] = True
    return mudou


def dados_da_om(oms: dict[str, dict[str, str]], nome: str) -> dict[str, str] | None:
    """OM salva com esse nome (sem diferenciar maiúsculas/acentos), ou None."""
    return oms.get(chave_om(nome))


def adicionar_dias_uteis(inicio: dt.date, dias: int) -> dt.date:
    """Data limite contando só segundas a sextas (feriados não são considerados)."""
    data = inicio
    while dias > 0:
        data += dt.timedelta(days=1)
        if data.weekday() < 5:
            dias -= 1
    return data


def gerar_protocolo(agora: dt.datetime | None = None) -> str:
    agora = agora or dt.datetime.now()
    return f"CD-{agora:%Y%m%d-%H%M}"


# ---------- itens ----------

# ---------- CATMAT automático ----------

def melhores_catmats(descricao: str, catmat, limite: int = MAX_CATMAT_SUGERIDOS) -> list[dict[str, object]]:
    """Os `limite` códigos CATMAT mais próximos da descrição (códigos diferentes, do mais para o menos parecido), com a % de combinação,
    qualquer que seja ela. Quem decide se entra na proposta é LIMIAR_CATMAT."""
    from catmat_busca import _opcoes_material  # import tardio: depende do Streamlit

    opcoes = sorted(_opcoes_material(descricao, catmat), key=lambda o: (round(o["bruta"], 1), o.get("popularidade", 0)), reverse=True)
    escolhidas, vistos = [], set()
    for opcao in opcoes:
        if opcao["codigo"] in vistos:
            continue
        vistos.add(opcao["codigo"])
        escolhidas.append({"codigo": str(opcao["codigo"]), "descricao": str(opcao["descricao_catalogo"]), "combinacao": float(opcao["similaridade"]),
                           "pdm": str(opcao["codigo_pdm"])})
        if len(escolhidas) >= limite:
            break
    return escolhidas


def sugerir_catmats(descricao: str, catmat, limite: int = MAX_CATMAT_SUGERIDOS, limiar: float = LIMIAR_CATMAT) -> list[dict[str, object]]:
    """Até `limite` códigos CATMAT com combinação >= `limiar`%. Vazio se nenhum chegar lá."""
    return [c for c in melhores_catmats(descricao, catmat, limite) if c["combinacao"] >= limiar]


def linha_da_tabela(item: dict[str, object], sugestoes: list[dict[str, object]], usar: bool = True) -> dict[str, object]:
    """Uma linha da tabela de itens da tela: o item e o CATMAT mais próximo (código, % de combinação e descrição do catálogo)."""
    melhor = sugestoes[0] if sugestoes else None
    entra = bool(melhor and melhor["combinacao"] >= LIMIAR_CATMAT)
    return {
        "Usar": usar, "Descrição": item["descricao"], "Qtd.": float(item["quantidade"]), "Un.": item["unidade"],
        "CATMAT": melhor["codigo"] if melhor else "—", "% casamento": round(melhor["combinacao"], 1) if melhor else 0.0,
        "Descrição do CATMAT": melhor["descricao"] if melhor else "Nenhuma correspondência no catálogo",
        "Outras opções": "\n".join(f"{c['codigo']} ({c['combinacao']:.0f}%)" for c in sugestoes[1:] if c["combinacao"] >= LIMIAR_CATMAT),
        "Situação": ("✅ entra na proposta" if entra else f"⚠️ abaixo de {LIMIAR_CATMAT:.0f}%: fica sem CATMAT (melhore a descrição)"),
    }


# ---------- fornecedores ----------

def _data_iso(valor: object) -> str:
    return str(valor or "")[:10]


def _data_br(valor: object) -> str:
    try:
        return dt.date.fromisoformat(_data_iso(valor)).strftime("%d/%m/%Y")
    except ValueError:
        return ""


def candidatos_do_item(descricao: str, catmat, hoje: dt.date | None = None, max_familias: int = 3) -> list[dict[str, object]]:
    """Fornecedores que venderam à Administração no último ano itens da mesma família (PDM), em ordem de preferência:
    1º os que venderam item muito parecido (combinação >= LIMIAR_CATMAT), do mais recente; depois os da mesma família (itens similares/mesma natureza)."""
    from catmat_busca import buscar_familias, calcular_similaridade
    from cotacao_rapida import _preco_valido, registros_da_familia

    hoje = hoje or dt.date.today()
    inicio, fim = (hoje - dt.timedelta(days=JANELA_DIAS)).isoformat(), hoje.isoformat()
    familias = [f for f in buscar_familias(descricao, catmat, limite=max_familias) if f["nota"] >= 50]
    por_cnpj: dict[str, dict[str, object]] = {}
    ufs_de_venda: dict[str, set[str]] = {}  # estados dos órgãos compradores a quem cada fornecedor vendeu (para preferir quem já vendeu no estado da OM)
    for familia in familias:
        registros, _ = registros_da_familia(str(familia["codigo"]), inicio, fim)
        for registro in registros:
            cnpj = somente_digitos(registro.get("cnpj"))
            if len(cnpj) != 14 or not _preco_valido(registro):
                continue
            ufs_de_venda.setdefault(cnpj, set()).add(str(registro.get("uf") or "").strip().upper())
            nota = calcular_similaridade(descricao, str(registro.get("descricao") or ""), str(familia["nome"]))
            categoria = 0 if nota >= LIMIAR_CATMAT else 1
            chave = (categoria, -dt.date.fromisoformat(_data_iso(registro["data"])).toordinal() if _data_iso(registro.get("data")) else 0)
            atual = por_cnpj.get(cnpj)
            if atual is None or chave < atual["chave"]:
                por_cnpj[cnpj] = {"chave": chave, "cnpj": cnpj, "nome": str(registro.get("fornecedor") or "").strip(), "parecido": categoria == 0,
                                  "data": _data_iso(registro.get("data")), "catmat": registro.get("catmat"), "descricao_vendida": str(registro.get("descricao") or ""),
                                  "uasg": str(registro.get("nome_uasg") or registro.get("uasg") or ""), "preco": registro.get("preco"), "familia": str(familia["nome"]),
                                  "combinacao": round(nota, 1)}
    for cnpj, dados in por_cnpj.items():
        dados["ufs_de_venda"] = ufs_de_venda.get(cnpj, set())
    return [dados for dados in sorted(por_cnpj.values(), key=lambda d: d["chave"])]


def normalizar_lugar(texto: object) -> str:
    """'São Paulo' = 'SAO PAULO' (sem acento, maiúsculas, espaços únicos)."""
    return chave_om(str(texto or ""))


def pontos_de_local(candidato: dict[str, object], contato: dict[str, str], local: dict[str, str] | None) -> int:
    """Quão perto do local da OM o fornecedor está: 3 = mesma cidade; 2 = mesmo estado; 1 = já vendeu a órgãos do mesmo estado; 0 = nenhum."""
    if not local:
        return 0
    uf_om, cidade_om = normalizar_lugar(local.get("uf")), normalizar_lugar(local.get("cidade"))
    if cidade_om and normalizar_lugar(contato.get("municipio")) == cidade_om and (not uf_om or normalizar_lugar(contato.get("uf")) in ("", uf_om)):
        return 3
    if uf_om and normalizar_lugar(contato.get("uf")) == uf_om:
        return 2
    if uf_om and uf_om in {normalizar_lugar(u) for u in candidato.get("ufs_de_venda", set())}:
        return 1
    return 0


def texto_do_local(pontos: int, local: dict[str, str] | None) -> str:
    if not local or not pontos:
        return ""
    return {3: f" Fica na mesma cidade da OM ({local.get('cidade', '')}/{local.get('uf', '')}).", 2: f" Fica no mesmo estado da OM ({local.get('uf', '')}).",
            1: f" Já vendeu a órgãos do mesmo estado da OM ({local.get('uf', '')})."}[pontos]


def motivo_da_escolha(c: dict[str, object]) -> str:
    """Justificativa da escolha do fornecedor (vai para o comprovante)."""
    data = _data_br(c.get("data"))
    onde = f" à {c['uasg']}" if c.get("uasg") else " à Administração Pública"
    if c.get("parecido"):
        return f"Vendeu item semelhante (CATMAT {c.get('catmat')}: {str(c.get('descricao_vendida'))[:80]}){onde} em {data} (compras.gov.br)."
    return f"Vendeu item da mesma família ({c.get('familia')}; CATMAT {c.get('catmat')}){onde} em {data} (compras.gov.br)."


def contato_do_cnpj(cnpj: str) -> dict[str, str]:
    """E-mail/telefone/razão social do CNPJ: OpenCNPJ e, se faltar e-mail, BrasilAPI. Campos vazios se não achar."""
    import fornecedores_nf

    dados = fornecedores_nf.interpretar_dados(fornecedores_nf.consultar_cnpj(cnpj, tentativas=2))
    email = dados.get("email", "").lower()
    contato = {"email": email if EMAIL_VALIDO.match(email) else "", "telefone": "" if dados.get("telefones", "").lower().startswith("n") else dados.get("telefones", ""),
               "razao_social": dados.get("razao_social", ""), "uf": dados.get("uf", ""), "municipio": dados.get("municipio", ""), "situacao": dados.get("situacao", "")}
    if not contato["email"]:
        try:
            resposta = requests.get(f"https://brasilapi.com.br/api/cnpj/v1/{somente_digitos(cnpj)}", timeout=8)
            if resposta.status_code == 200:
                extra = resposta.json()
                email = str(extra.get("email") or "").strip().lower()
                if EMAIL_VALIDO.match(email):
                    contato["email"] = email
                contato["razao_social"] = contato["razao_social"] or str(extra.get("razao_social") or "")
                contato["uf"] = contato["uf"] or str(extra.get("uf") or "")
                contato["situacao"] = contato["situacao"] or str(extra.get("descricao_situacao_cadastral") or "")
                contato["telefone"] = contato["telefone"] or str(extra.get("ddd_telefone_1") or "").strip()
        except (requests.RequestException, ValueError):
            pass
    return contato


def _ativa(contato: dict[str, str]) -> bool:
    situacao = contato.get("situacao", "").lower()
    return not any(marca in situacao for marca in SITUACOES_INATIVAS)


def _candidatos_seguro(descricao: str, catmat) -> list[dict[str, object]]:
    try:
        return candidatos_do_item(descricao, catmat)
    except Exception:  # falha de rede em um item não derruba o pedido inteiro
        return []


def escolher_fornecedores(itens: list[dict[str, object]], catmat, ao_progredir=None, excluir_emails: set[str] | None = None,
                          por_item: int = FORNECEDORES_POR_ITEM, local: dict[str, str] | None = None) -> dict[int, list[dict[str, object]]]:
    """Para cada item (na ordem), até `por_item` fornecedores COM e-mail, nunca repetindo um fornecedor (CNPJ ou e-mail) já escolhido para outro item.
    Devolve {posição do item: [fornecedor, ...]}.
    Etapa 1 (em paralelo): vendas recentes de cada item no Compras.gov. Etapa 2 (item a item, para não repetir fornecedor): e-mail dos candidatos.
    `local` ({cidade, uf} da OM): entre fornecedores igualmente parecidos com o item, dá preferência aos da mesma cidade, depois do mesmo estado e depois a quem já
    vendeu a órgãos do mesmo estado (a localização do fornecedor vem do cadastro do CNPJ, por isso são consultados alguns candidatos a mais).
    `ao_progredir(fase, feitos, total, texto, achados)` atualiza a tela (fase 'busca' ou 'contatos')."""
    total = len(itens)
    candidatos_por_item: dict[int, list[dict[str, object]]] = {}
    with ThreadPoolExecutor(max_workers=4) as executor:
        futuros = {executor.submit(_candidatos_seguro, str(item["descricao"]), catmat): posicao for posicao, item in enumerate(itens)}
        for feitos, futuro in enumerate(as_completed(futuros), start=1):
            candidatos_por_item[futuros[futuro]] = futuro.result()
            if ao_progredir:
                ao_progredir("busca", feitos, total, f"Vendas recentes consultadas: {feitos}/{total} itens", 0)
    usados_cnpj: set[str] = set()
    usados_email: set[str] = {e.lower() for e in (excluir_emails or set())}
    resultado: dict[int, list[dict[str, object]]] = {}
    achados = 0
    for posicao, item in enumerate(itens):
        if ao_progredir:
            ao_progredir("contatos", posicao, total, f"Item {posicao + 1}/{total}: {item['descricao']}", achados)
        candidatos = [c for c in candidatos_por_item.get(posicao, []) if c["cnpj"] not in usados_cnpj]
        if local and local.get("uf"):  # quem já vendeu no estado da OM vem antes, dentro da mesma categoria (parecido x mesma família)
            uf_om = normalizar_lugar(local["uf"])
            candidatos.sort(key=lambda c: (c["chave"][0], 0 if uf_om in {normalizar_lugar(u) for u in c.get("ufs_de_venda", set())} else 1, c["chave"][1]))
        validos: list[tuple[tuple[int, int, int], dict[str, object], dict[str, str]]] = []  # (chave de ordem, candidato, contato)
        consultados = 0
        meta = por_item * (3 if local else 1)  # com preferência de local, junta mais candidatos com e-mail para escolher entre eles
        ordem = 0
        while candidatos and len(validos) < meta and consultados < MAX_CONSULTAS_CONTATO_POR_ITEM:
            falta = meta - len(validos)
            lote, candidatos = candidatos[:max(falta, 4)], candidatos[max(falta, 4):]
            consultados += len(lote)
            with ThreadPoolExecutor(max_workers=4) as executor:
                contatos = list(executor.map(lambda c: contato_do_cnpj(str(c["cnpj"])), lote))
            for candidato, contato in zip(lote, contatos):
                email = contato.get("email", "").lower()
                ordem += 1
                if not email or not _ativa(contato) or email in usados_email:
                    continue
                validos.append(((candidato["chave"][0], -pontos_de_local(candidato, contato, local), ordem), candidato, contato))
            time.sleep(0.1)
        escolhidos: list[dict[str, object]] = []
        for _, candidato, contato in sorted(validos, key=lambda v: v[0]):
            email = contato.get("email", "").lower()
            if len(escolhidos) >= por_item:
                break
            if email in usados_email:  # dois candidatos com o mesmo e-mail
                continue
            usados_cnpj.add(str(candidato["cnpj"]))
            usados_email.add(email)
            pontos = pontos_de_local(candidato, contato, local)
            escolhidos.append({**candidato, **contato, "email": email, "nome": contato.get("razao_social") or candidato["nome"], "pontos_local": pontos,
                               "motivo": motivo_da_escolha(candidato) + texto_do_local(pontos, local), "item": item["descricao"], "posicao_item": posicao + 1, "enviar": True})
        resultado[posicao] = escolhidos
        achados += len(escolhidos)
    if ao_progredir:
        ao_progredir("contatos", total, total, "Fornecedores escolhidos", achados)
    return resultado


# ---------- texto do e-mail ----------

def assunto_email(protocolo: str, nome_om: str) -> str:
    return f"Solicitação de Cotação nº {protocolo} – {nome_om}"


def texto_email(om: dict[str, str], protocolo: str, data_limite: dt.date, prazo_dias: int, validade_dias: int, total_itens: int, saudacao: str = "Prezados Senhores,") -> str:
    contato = " | ".join(p for p in (om.get("telefone", ""), om.get("email", "")) if p)
    return (
        f"{saudacao}\n\n"
        f"A Organização Militar {om['nome']} vem, respeitosamente, convidar a sua empresa a apresentar proposta de preços para o fornecimento dos {total_itens} item(ns) "
        f"relacionados na planilha em anexo (Solicitação de Cotação nº {protocolo}).\n\n"
        "Sua empresa foi identificada por ter fornecido, recentemente, materiais iguais ou semelhantes a órgãos da Administração Pública Federal. "
        "Por isso, sua cotação é muito valiosa para que a nossa pesquisa de preços reflita a realidade do mercado — e, quem sabe, para abrir uma "
        "nova oportunidade de negócio com a Marinha do Brasil.\n\n"
        "Para participar, pedimos a gentileza de:\n"
        "  1. Preencher o arquivo em anexo (Word, editável), com os valores unitários dos itens que a empresa puder fornecer — não é necessário cotar todos;\n"
        "  2. Completar os dados da empresa (razão social, CNPJ, endereço, telefone e e-mail), a data e a assinatura do responsável;\n"
        f"  3. Responder a este e-mail, anexando o arquivo preenchido, até {data_limite:%d/%m/%Y} ({prazo_dias} dias úteis). "
        f"A proposta deve ter validade mínima de {validade_dias} dias.\n\n"
        "Esclarecemos que este pedido tem finalidade exclusiva de pesquisa de preços (Lei nº 14.133/2021, art. 23, e IN SEGES/ME nº 65/2021) "
        "e não representa compromisso de contratação. Se a empresa não tiver interesse, agradecemos o retorno com essa informação.\n\n"
        "Em caso de dúvidas, estamos à disposição"
        + (f": {contato}.\n\n" if contato else ".\n\n")
        + "Agradecemos desde já a atenção e a colaboração.\n\n"
        f"Atenciosamente,\n{om.get('responsavel') or 'Setor responsável pela pesquisa de preços'}\n{om['nome']}"
        + (f"\nCNPJ {formatar_cnpj(om['cnpj'])}" if om.get("cnpj") else "") + f"\n{endereco_completo(om)}"
    )


def html_email(texto: str) -> str:
    """Mesmo texto, em HTML simples (parágrafos e a lista numerada)."""
    blocos = []
    for bloco in texto.split("\n\n"):
        linhas = [html.escape(l) for l in bloco.split("\n")]
        blocos.append("<p style=\"margin:0 0 14px 0\">" + "<br>".join(l.replace("  ", "&nbsp;&nbsp;") for l in linhas) + "</p>")
    return ("<div style=\"font-family:Arial,Helvetica,sans-serif;font-size:14px;line-height:1.5;color:#1a1a1a;max-width:680px\">"
            + "".join(blocos) + "</div>")


# ---------- proposta em Word ----------

def _sombrear(celula, cor: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    propriedades = celula._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor)
    propriedades.append(sombra)


def _bordas(tabela, cor: str = "9AA5B1", tamanho: int = 4) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    propriedades = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elemento = OxmlElement(f"w:{lado}")
        elemento.set(qn("w:val"), "single")
        elemento.set(qn("w:sz"), str(tamanho))
        elemento.set(qn("w:space"), "0")
        elemento.set(qn("w:color"), cor)
        bordas.append(elemento)
    propriedades.append(bordas)


def _fixar_larguras(tabela, larguras) -> None:
    """Larguras de coluna que o Word e o LibreOffice respeitam (layout fixo + largura em cada célula)."""
    tabela.autofit = False
    for coluna, largura in enumerate(larguras):
        tabela.columns[coluna].width = largura
        for celula in tabela.columns[coluna].cells:
            celula.width = largura


def _repetir_cabecalho(linha) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    propriedades = linha._tr.get_or_add_trPr()
    elemento = OxmlElement("w:tblHeader")
    elemento.set(qn("w:val"), "true")
    propriedades.append(elemento)


def _campo_pagina(paragrafo, tamanho) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def run_com(tipo=None, texto=None):
        run = paragrafo.add_run()
        run.font.size = tamanho
        if tipo:
            elemento = OxmlElement("w:fldChar")
            elemento.set(qn("w:fldCharType"), tipo)
            run._r.append(elemento)
        if texto:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = texto
            run._r.append(instr)

    run_com("begin")
    run_com(texto=" PAGE ")
    run_com("end")


def gerar_proposta_docx(om: dict[str, str], itens: list[dict[str, object]], catmats: dict[int, list[dict[str, object]]], protocolo: str,
                        data_emissao: dt.date, data_limite: dt.date, prazo_dias: int, validade_dias: int) -> bytes:
    """Proposta de preços editável (Word, paisagem): cabeçalho da OM, dados da empresa, condições, tabela de itens, declaração e assinatura."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    azul, dourado, claro = RGBColor(0x00, 0x1A, 0x4D), RGBColor(0xB8, 0x92, 0x1F), "EAF0FA"
    doc = Document()
    secao = doc.sections[0]
    secao.orientation = WD_ORIENT.LANDSCAPE
    secao.page_width, secao.page_height = Cm(29.7), Cm(21.0)
    secao.left_margin = secao.right_margin = Cm(1.6)
    secao.top_margin, secao.bottom_margin = Cm(3.0), Cm(1.8)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10)
    estilo.paragraph_format.space_after = Pt(2)

    # cabeçalho de página: dados da OM
    cabecalho = secao.header
    paragrafo = cabecalho.paragraphs[0]
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragrafo.add_run("MARINHA DO BRASIL\n")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(9), azul
    run = paragrafo.add_run(om["nome"].upper())
    run.bold, run.font.size, run.font.color.rgb = True, Pt(13), azul
    detalhe = cabecalho.add_paragraph()
    detalhe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = detalhe.add_run(" | ".join(p for p in (f"CNPJ {formatar_cnpj(om['cnpj'])}" if om.get("cnpj") else "", endereco_completo(om), om.get("telefone", ""),
                                                  om.get("email", "")) if p))
    run.font.size, run.font.color.rgb = Pt(8), RGBColor(0x55, 0x5F, 0x6B)
    linha = cabecalho.add_paragraph()
    run = linha.add_run("━" * 90)
    run.font.size, run.font.color.rgb = Pt(6), dourado
    linha.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # rodapé
    rodape = secao.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = rodape.add_run(f"Solicitação de Cotação nº {protocolo}  •  Página ")
    run.font.size = Pt(8)
    _campo_pagina(rodape, Pt(8))

    # título
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("SOLICITAÇÃO DE COTAÇÃO – PROPOSTA DE PREÇOS")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(16), azul
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f"Nº {protocolo}   •   Emitida em {data_emissao:%d/%m/%Y}   •   Resposta até {data_limite:%d/%m/%Y} ({prazo_dias} dias úteis)")
    run.font.size, run.font.color.rgb = Pt(10), RGBColor(0x33, 0x3D, 0x4A)

    def subtitulo(texto: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(texto)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(10.5), azul

    def tabela_campos(campos: list[tuple[str, str]], colunas: int = 2) -> None:
        linhas = (len(campos) + colunas - 1) // colunas
        tabela = doc.add_table(rows=linhas, cols=colunas * 2)
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordas(tabela)
        for indice, (rotulo, valor) in enumerate(campos):
            linha_, coluna = indice // colunas, (indice % colunas) * 2
            c_rotulo, c_valor = tabela.cell(linha_, coluna), tabela.cell(linha_, coluna + 1)
            _sombrear(c_rotulo, claro)
            c_rotulo.paragraphs[0].add_run(rotulo).bold = True
            c_valor.paragraphs[0].add_run(valor)
            c_rotulo.width, c_valor.width = Cm(4.2), Cm(9.2)
        for linha_ in tabela.rows:
            linha_.height = Cm(0.75)
        _fixar_larguras(tabela, [Cm(4.2), Cm(9.2)] * colunas)
        sobra = len(campos) % colunas
        if sobra:  # último campo sozinho na linha: o espaço de resposta ocupa a largura toda
            ultima = tabela.rows[-1]
            ultima.cells[sobra * 2 - 1].merge(ultima.cells[-1])

    subtitulo("1. DADOS DA EMPRESA PROPONENTE (a preencher)")
    tabela_campos([("Razão social", ""), ("CNPJ", ""), ("Endereço", ""), ("Cidade / UF / CEP", ""), ("Telefone", ""), ("E-mail", ""),
                   ("Responsável (nome e cargo)", ""), ("Inscrição estadual / municipal", "")])
    subtitulo("2. CONDIÇÕES DA PROPOSTA (a preencher)")
    tabela_campos([("Validade da proposta", f"____ dias (mínimo {validade_dias} dias)"), ("Prazo de entrega", "____ dias corridos após o pedido"),
                   ("Condições de pagamento", ""), ("Garantia", ""), ("Observações", "")])

    subtitulo("3. ITENS COTADOS")
    cabecalhos = ["ITEM", "CATMAT\n(sugerido)", "DESCRIÇÃO", "UN.", "QTD.", "MARCA / MODELO\n(a preencher)", "VALOR UNIT.\n(R$)", "VALOR TOTAL\n(R$)"]
    larguras = [Cm(1.2), Cm(2.6), Cm(9.2), Cm(1.4), Cm(1.6), Cm(4.4), Cm(3.0), Cm(3.0)]
    tabela = doc.add_table(rows=1, cols=len(cabecalhos))
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordas(tabela)
    for coluna, texto in enumerate(cabecalhos):
        celula = tabela.rows[0].cells[coluna]
        _sombrear(celula, "001A4D")
        p = celula.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0xFF, 0xFF, 0xFF)
        celula.width = larguras[coluna]
    _repetir_cabecalho(tabela.rows[0])
    for posicao, item in enumerate(itens):
        linha_ = tabela.add_row()
        codigos = "\n".join(c["codigo"] for c in catmats.get(posicao, []))
        valores = [str(posicao + 1), codigos, str(item["descricao"]), str(item["unidade"]), formatar_quantidade(item["quantidade"]), "", "", ""]
        for coluna, valor in enumerate(valores):
            celula = linha_.cells[coluna]
            celula.width = larguras[coluna]
            if posicao % 2:
                _sombrear(celula, "F4F7FB")
            p = celula.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if coluna in (2, 5) else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(valor)
            r.font.size = Pt(9)
        linha_.height = Cm(0.9)
    _fixar_larguras(tabela, larguras)
    total = tabela.add_row()
    fundida = total.cells[0].merge(total.cells[6])
    fundida.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fundida.paragraphs[0].add_run("VALOR TOTAL GLOBAL (R$)").bold = True
    _sombrear(fundida, claro)
    _sombrear(total.cells[7], claro)

    nota = doc.add_paragraph()
    nota.paragraph_format.space_before = Pt(4)
    run = nota.add_run("O código CATMAT é apenas uma referência para facilitar a identificação do material; em caso de dúvida sobre a descrição, a empresa pode "
                       "indicar a marca/modelo ofertado e eventuais observações. Não é necessário cotar todos os itens.")
    run.italic, run.font.size = True, Pt(8)

    subtitulo("4. DECLARAÇÃO")
    declaracao = doc.add_paragraph("Declaramos que os preços acima incluem todos os custos e despesas (tributos, frete, embalagem, seguro e demais encargos) necessários ao "
                                   "fornecimento, e que esta proposta tem finalidade de pesquisa de preços, sem caracterizar compromisso de contratação.")
    declaracao.paragraph_format.space_after = Pt(14)

    assinatura = doc.add_table(rows=2, cols=2)
    assinatura.alignment = WD_TABLE_ALIGNMENT.CENTER
    assinatura.cell(0, 0).paragraphs[0].add_run("_____________________, ____ de ______________ de ________.").font.size = Pt(10)
    p = assinatura.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("\n\n_____________________________________________")
    p = assinatura.cell(1, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Assinatura e carimbo do responsável pela empresa")
    r.font.size = Pt(8.5)

    memoria = io.BytesIO()
    doc.save(memoria)
    return memoria.getvalue()


# ---------- comprovante (Excel) ----------

def gerar_comprovante_xlsx(om: dict[str, str], protocolo: str, itens: list[dict[str, object]], catmats: dict[int, list[dict[str, object]]],
                           fornecedores: list[dict[str, object]], envios: list[dict[str, object]] | None = None) -> bytes:
    """Planilha de registro: itens com CATMAT, fornecedores escolhidos com a justificativa e o resultado do envio."""
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    linhas_itens = [{"Item": i + 1, "Descrição": it["descricao"], "Un.": it["unidade"], "Qtd.": formatar_quantidade(it["quantidade"]),
                     "CATMAT sugerido": " / ".join(f"{c['codigo']} ({c['combinacao']:.0f}%)" for c in catmats.get(i, [])) or "—",
                     "Fornecedores com e-mail": sum(1 for f in fornecedores if f["posicao_item"] == i + 1)} for i, it in enumerate(itens)]
    linhas_forn = [{"Item": f["posicao_item"], "Descrição do item": f["item"], "Empresa": f["nome"], "CNPJ": f["cnpj"], "E-mail": f["email"],
                    "Telefone": f.get("telefone", ""), "Cidade/UF": f.get("cidade_uf") or "/".join(p for p in (str(f.get("municipio", "")).title(), str(f.get("uf", ""))) if p), "Enviado": "Sim" if f.get("enviar", True) else "Não",
                    "Motivo da escolha": f["motivo"]} for f in fornecedores]
    resumo = [{"Campo": "Protocolo", "Valor": protocolo}, {"Campo": "OM", "Valor": om["nome"]}, {"Campo": "CNPJ da OM", "Valor": formatar_cnpj(om.get("cnpj", ""))}, {"Campo": "Endereço", "Valor": endereco_completo(om)},
              {"Campo": "E-mail da OM", "Valor": om.get("email", "")}, {"Campo": "Emitido em", "Valor": dt.datetime.now().strftime("%d/%m/%Y %H:%M")},
              {"Campo": "Fornecedores (e-mails distintos)", "Valor": len({f['email'] for f in fornecedores if f.get('enviar', True)})}]
    saida = io.BytesIO()
    with pd.ExcelWriter(saida, engine="openpyxl") as escritor:
        pd.DataFrame(resumo).to_excel(escritor, sheet_name="Resumo", index=False)
        pd.DataFrame(linhas_itens).to_excel(escritor, sheet_name="Itens e CATMAT", index=False)
        pd.DataFrame(linhas_forn).to_excel(escritor, sheet_name="Fornecedores", index=False)
        if envios:
            pd.DataFrame(envios).to_excel(escritor, sheet_name="Envio", index=False)
        for aba in escritor.book.worksheets:
            for celula in aba[1]:
                celula.font = Font(bold=True, color="FFFFFF")
                celula.fill = PatternFill("solid", fgColor="001A4D")
            for coluna in aba.columns:
                maior = max(len(str(c.value or "")) for c in coluna)
                aba.column_dimensions[get_column_letter(coluna[0].column)].width = min(max(12, maior + 2), 70)
                for c in coluna:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    return saida.getvalue()


# ---------- envio (Resend) ----------

def remetente_com_nome(remetente: str, nome_om: str) -> str:
    """'Nome <a@b.com>' ou 'a@b.com' -> 'Cotação – NOME DA OM <a@b.com>'."""
    casamento = re.search(r"<([^>]+)>", remetente)
    endereco = (casamento.group(1) if casamento else remetente).strip()
    nome = re.sub(r"[<>\"]", "", f"Cotação – {nome_om}")[:80]
    return f"{nome} <{endereco}>"


def senha_correta(informada: str, esperada: str) -> bool:
    return bool(esperada) and hmac.compare_digest(str(informada).encode(), str(esperada).encode())


def lista_emails(texto: str) -> tuple[list[str], list[str]]:
    """'a@x.com; b@y.com' -> (válidos sem repetir, trechos inválidos). Separa por ponto e vírgula, vírgula ou espaço."""
    validos: list[str] = []
    invalidos: list[str] = []
    for parte in re.split(r"[;,\s]+", str(texto or "")):
        parte = parte.strip().lower()
        if not parte:
            continue
        if EMAIL_VALIDO.match(parte):
            if parte not in validos:
                validos.append(parte)
        else:
            invalidos.append(parte)
    return validos, invalidos


def separar_destinatarios(emails_om: list[str] | str, fornecedores: list[str], lote: int = LOTE_BCC) -> list[list[str]]:
    """E-mails ocultos em lotes, sem repetir (sem diferenciar maiúsculas) e sem incluir os e-mails da OM."""
    vistos = {e.strip().lower() for e in ([emails_om] if isinstance(emails_om, str) else emails_om)}
    unicos = []
    for email in fornecedores:
        chave = email.strip().lower()
        if chave and chave not in vistos and EMAIL_VALIDO.match(chave):
            vistos.add(chave)
            unicos.append(chave)
    return [unicos[i:i + lote] for i in range(0, len(unicos), lote)]


def enviar_resend(chave_api: str, remetente: str, para: list[str] | str, bcc: list[str], responder_para: list[str] | str, assunto: str, texto: str, html_corpo: str,
                  anexos: list[tuple[str, bytes]], sessao=None) -> dict[str, object]:
    """Uma mensagem pela API do Resend (`para` e `responder_para` aceitam vários e-mails). Devolve {'ok', 'id', 'erro'}. Nunca levanta exceção."""
    corpo: dict[str, object] = {"from": remetente, "to": [para] if isinstance(para, str) else list(para), "subject": assunto, "text": texto, "html": html_corpo,
                                "reply_to": responder_para if isinstance(responder_para, str) else list(responder_para)}
    if bcc:
        corpo["bcc"] = bcc
    if anexos:
        corpo["attachments"] = [{"filename": nome, "content": base64.b64encode(conteudo).decode("ascii")} for nome, conteudo in anexos]
    try:
        resposta = (sessao or requests).post(API_RESEND, json=corpo, headers={"Authorization": f"Bearer {chave_api}", "Content-Type": "application/json"}, timeout=60)
    except requests.RequestException as erro:
        return {"ok": False, "id": "", "erro": f"falha de rede ({type(erro).__name__})"}
    if resposta.status_code in (200, 201):
        try:
            return {"ok": True, "id": resposta.json().get("id", ""), "erro": ""}
        except ValueError:
            return {"ok": True, "id": "", "erro": ""}
    return {"ok": False, "id": "", "erro": explicar_erro_resend(resposta.status_code, resposta.text)}


def explicar_erro_resend(status: int, texto: str) -> str:
    """Erros comuns do Resend em português, com o que fazer; senão, o código e o começo da resposta."""
    minusculo = texto.lower()
    if status == 403 and "only send testing emails" in minusculo:
        return ("O remetente de teste do Resend (onboarding@resend.dev) só envia para o e-mail da conta do Resend. Para enviar a outros destinatários, "
                "verifique um domínio no Resend (Domains) e use um endereço dele em RESEND_FROM nos Secrets. Para só testar, ponha o e-mail da conta do Resend no campo 'Seu e-mail institucional'.")
    if status == 403 and "domain is not verified" in minusculo:
        return "O domínio do remetente (RESEND_FROM) ainda não está verificado no Resend. Conclua a verificação em Domains e tente de novo."
    if status in (401, 403) and ("api key" in minusculo or "unauthorized" in minusculo):
        return "A chave do Resend (RESEND_API_KEY) não foi aceita. Confira se foi copiada inteira nos Secrets, sem espaços."
    if status == 429:
        return "O Resend limitou os envios (limite por segundo ou por dia do plano). Aguarde um pouco e tente de novo."
    return f"HTTP {status}: {texto[:160]}"
