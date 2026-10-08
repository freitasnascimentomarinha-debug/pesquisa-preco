"""Relatórios (PDF e Excel) da Cotação Rápida, no layout visual do relatório da página Cotação."""

from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from fpdf import FPDF
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lista_itens import formatar_quantidade, tem_quantidades, valor_total_item, valor_total_orcamento
from cotacao_rapida import JANELA_DIAS, LIMIAR_CORRESPONDENCIA, LIMIAR_SERVICO, MAX_CATMAT, MAX_PRECOS, MIN_PRECOS, TOLERANCIA

AZUL, DOURADO, AZUL_CARTAO = (0, 26, 77), (212, 175, 55), (10, 37, 64)
LARGURA = 281  # área útil (A4 paisagem, margens de 8 mm)

JUSTIFICATIVA_COTACAO = (
    "Justificativa quanto aos critérios de Pesquisa de Preços (IN 65/2021)\n\n"
    "A presente pesquisa de preços foi realizada em estrita observância aos critérios e parâmetros "
    "estabelecidos pela Instrução Normativa SEGES/ME nº 65, de 7 de julho de 2021. Foram priorizados "
    "dados provenientes de aquisições e contratações similares no âmbito da Administração Pública, "
    "assegurando a contemporaneidade dos registros e a compatibilidade técnica com o objeto pretendido.\n\n"
    "Os valores obtidos refletem os preços efetivamente praticados no mercado, garantindo a "
    "economicidade e a ampla fundamentação na formação do preço estimado para a contratação. "
    "Ressalta-se que a utilização de painéis de preços e bancos de dados oficiais confere "
    "transparência, rastreabilidade e segurança documental ao procedimento de aferição do valor de mercado.\n\n"
    "Dessa forma, entende-se que a metodologia adotada atende aos princípios da razoabilidade, "
    "economicidade e motivação do ato administrativo, conferindo robustez à formação do preço estimado."
)
def metodologia() -> str:
    return (
    "Metodologia da Cotação Rápida\n\n"
    f"1. Para cada descrição informada, foram localizados até {MAX_CATMAT} itens do catálogo CATMAT (materiais) ou CATSERV (serviços) "
    f"com correspondência igual ou superior a {LIMIAR_CORRESPONDENCIA:.0f}% (materiais) ou {LIMIAR_SERVICO:.0f}% (serviços) e que possuem preços praticados no período.\n\n"
    f"2. Os preços foram obtidos no módulo Pesquisa de Preço do Compras.gov (dados abertos), considerando compras dos "
    f"últimos {JANELA_DIAS} dias, e comparados apenas dentro da mesma unidade de fornecimento (materiais) ou de medida (serviços).\n\n"
    "3. Foram descartados preços inexequíveis ou extremos (abaixo de 30% ou acima de 300% da mediana) e os valores atípicos "
    "(outliers) pelo método do intervalo interquartil (IQR).\n\n"
    f"4. Foram selecionados até {MAX_PRECOS} preços, priorizando fornecedores distintos e as compras mais recentes, de modo que "
    f"cada preço listado não se afaste mais que {TOLERANCIA * 100:.0f}% (para mais ou para menos) da média dos preços listados. "
    f"Itens com menos de {MIN_PRECOS} preços são sinalizados e devem ser complementados por outras fontes (IN SEGES/ME nº 65/2021, art. 5º).\n\n"
    "5. A média unitária apresentada é a média aritmética dos preços listados. A correspondência entre a descrição e o item CATMAT "
    "é uma sugestão automática e deve ser conferida pelo requisitante antes do uso no processo. Em serviços, o preço depende do escopo (quantidade, área, jornada), que deve ser compatível com o objeto pretendido."
)

STATUS_TEXTO = {
    "ok": "Preços suficientes",
    "insuficiente": f"Menos de {MIN_PRECOS} preços",
    "sem_precos": "Sem preços válidos",
    "sem_catmat": "Sem CATMAT/CATSERV com preços",
}


def _seguro(texto: object) -> str:
    if texto is None or str(texto) in ("nan", "None"):
        return ""
    return str(texto).replace("\n", " ").replace("\r", " ").encode("latin-1", "replace").decode("latin-1")


def moeda(valor: float | None) -> str:
    if valor is None or valor != valor:
        return ""
    return f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _quebrar(pdf: FPDF, texto: str, largura: float, max_linhas: int) -> list[str]:
    """Quebra o texto em linhas que cabem na largura (em mm), com reticências se passar do limite."""
    palavras, linhas, atual = _seguro(texto).split(), [], ""
    for palavra in palavras:
        teste = f"{atual} {palavra}".strip()
        if pdf.get_string_width(teste) <= largura or not atual:
            atual = teste
        else:
            linhas.append(atual)
            atual = palavra
    if atual:
        linhas.append(atual)
    if len(linhas) > max_linhas:
        linhas = linhas[:max_linhas]
        while linhas[-1] and pdf.get_string_width(linhas[-1] + "...") > largura:
            linhas[-1] = linhas[-1][:-1]
        linhas[-1] += "..."
    return linhas or [""]


def _truncar(pdf: FPDF, texto: object, largura: float) -> str:
    texto = _seguro(texto)
    while texto and pdf.get_string_width(texto) > largura:
        texto = texto[:-1]
    return texto


class PDFCotacaoRapida(FPDF):
    def __init__(self, *args, titulo: str = "Cotacao Rapida - Mapa Comparativo de Precos", **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.titulo = titulo

    def header(self) -> None:
        self.set_fill_color(*AZUL)
        self.rect(0, 0, self.w, 30, "F")
        self.set_fill_color(*DOURADO)
        self.rect(0, 30, self.w, 1.2, "F")
        self.set_font("Helvetica", "B", 20)
        self.set_text_color(*DOURADO)
        self.set_xy(10, 4)
        self.cell(60, 10, "AtaCotada")
        self.set_font("Helvetica", "B", 11)
        self.set_text_color(255, 255, 255)
        self.set_xy(10, 16)
        self.cell(0, 8, _seguro(self.titulo))
        self.set_font("Helvetica", "", 7)
        self.set_text_color(180, 180, 180)
        self.set_xy(self.w - 70, 6)
        self.cell(60, 5, "Marinha do Brasil", align="R")
        self.set_y(35)

    def footer(self) -> None:
        self.set_y(-10)
        self.set_fill_color(*AZUL)
        self.rect(0, self.h - 10, self.w, 10, "F")
        self.set_font("Helvetica", "I", 6.5)
        self.set_text_color(160, 160, 160)
        self.cell(0, 8, f"AtaCotada - Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", align="L")
        self.set_x(-30)
        self.cell(20, 8, f"{self.page_no()}/{{nb}}", align="R")


def _cartoes(pdf: FPDF, itens: list[tuple[str, str]], y: float, altura: float = 12) -> None:
    largura = LARGURA / len(itens)
    pdf.set_y(y)
    for rotulo, valor in itens:
        x, y0 = pdf.get_x(), pdf.get_y()
        pdf.set_fill_color(*AZUL_CARTAO)
        pdf.rect(x, y0, largura - 1, altura, "F")
        pdf.set_fill_color(*DOURADO)
        pdf.rect(x, y0, 1.2, altura, "F")
        pdf.set_xy(x + 2, y0 + 1.5)
        pdf.set_font("Helvetica", "", 6)
        pdf.set_text_color(*DOURADO)
        pdf.cell(largura - 4, 3, _seguro(rotulo))
        pdf.set_xy(x + 2, y0 + 5)
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_text_color(255, 255, 255)
        pdf.cell(largura - 4, 5, _seguro(valor))
        pdf.set_xy(x + largura, y0)
    pdf.ln(altura + 3)


def _cabecalho_tabela(pdf: FPDF, colunas: list[tuple[str, float]], altura: float = 7) -> None:
    pdf.set_fill_color(*AZUL)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 7)
    for rotulo, largura in colunas:
        pdf.cell(largura, altura, _seguro(rotulo), 1, 0, "C", True)
    pdf.ln()


COLUNAS_MAPA = [("#", 8), ("Item pesquisado", 66), ("Codigo (correspondencia)", 36), ("Unid.", 24)] + [
    (f"Preco {n}", 19) for n in range(1, MAX_PRECOS + 1)
] + [("Media unitaria", 26), ("N. precos", 11), ("CV", 13)]


def _linhas_codigos(r: dict) -> list[str]:
    rotulo_codigo = "CATSERV" if r.get("tipo") == "Serviço" else "CATMAT"
    return [f"{rotulo_codigo} {k['codigo']} ({k['correspondencia']:.0f}%)" for k in r["catmats"]] or ["-"]


def _colunas_mapa(max_precos: int) -> list[tuple[str, float]]:
    """Colunas do mapa para `max_precos` colunas de preço; a largura que sobra (ou falta) vai para a coluna do item."""
    colunas = [(rotulo, largura) for rotulo, largura in COLUNAS_MAPA if not rotulo.startswith("Preco ")]
    colunas[1] = (colunas[1][0], colunas[1][1] + (MAX_PRECOS - max_precos) * 19)
    return colunas[:4] + [(f"Preco {n}", 19) for n in range(1, max_precos + 1)] + colunas[4:]


def _pagina_mapa(pdf: FPDF, resultados: list[dict], rotulo_col3: str = "Codigo (correspondencia)", linhas_col3=_linhas_codigos, fonte: str | None = None, max_precos: int = MAX_PRECOS) -> None:
    colunas_mapa = [(rotulo_col3, largura) if indice == 2 else (rotulo, largura) for indice, (rotulo, largura) in enumerate(_colunas_mapa(max_precos))]
    total = len(resultados)
    ok = sum(1 for r in resultados if r["status"] == "ok")
    insuficientes = sum(1 for r in resultados if r["status"] == "insuficiente")
    sem = total - ok - insuficientes
    _cartoes(pdf, [("ITENS PESQUISADOS", str(total)), (f"COM {MIN_PRECOS}+ PRECOS", str(ok)), (f"MENOS DE {MIN_PRECOS} PRECOS", str(insuficientes)), ("SEM PRECO VALIDO", str(sem))], 38)

    pdf.set_font("Helvetica", "B", 9)
    pdf.set_text_color(*AZUL)
    pdf.cell(0, 6, "MAPA COMPARATIVO DE PRECOS - precos unitarios praticados (R$), com a media unitaria de cada item", ln=True)
    pdf.set_font("Helvetica", "", 6.5)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 4, "Verde: preco igual ou abaixo da media do item  |  Ambar: acima da media  |  Cada preco fica a ate 30% da media do item.", ln=True)
    pdf.ln(1)
    _cabecalho_tabela(pdf, colunas_mapa)

    for numero, r in enumerate(resultados, start=1):
        pdf.set_font("Helvetica", "", 6.5)
        desc = _quebrar(pdf, r["descricao"], colunas_mapa[1][1] - 2, 3)
        catmats = linhas_col3(r)
        linhas = max(len(desc), len(catmats), 1)
        altura = max(8.0, 3.6 * linhas + 2)
        if pdf.get_y() + altura > pdf.h - 16:
            pdf.add_page()
            pdf.set_y(35)
            _cabecalho_tabela(pdf, colunas_mapa)
        x0, y0 = pdf.l_margin, pdf.get_y()
        precos = [p["preco"] for p in r["precos"]]
        media = r["stats"]["media"] if r["stats"] else None
        fundo = (245, 248, 255) if numero % 2 == 0 else (255, 255, 255)
        x = x0
        for indice, (_, largura) in enumerate(colunas_mapa):
            if not precos and 4 < indice < 4 + max_precos:
                continue  # as colunas de preço de um item sem preços viram uma só, com a situação
            if not precos and indice == 4:
                largura = sum(w for _, w in colunas_mapa[4:4 + max_precos])
            cor = fundo
            texto: str | list[str] = ""
            alinhamento = "C"
            fonte_negrito = False
            if indice == 0:
                texto = str(numero)
            elif indice == 1:
                texto, alinhamento = desc, "L"
            elif indice == 2:
                texto, alinhamento = catmats, "L"
            elif indice == 3:
                texto = _truncar(pdf, r.get("unidade_curta") or r["unidade"], largura - 2) if r["unidade"] else "-"
            elif 4 <= indice < 4 + max_precos:
                posicao = indice - 4
                if posicao < len(precos):
                    texto = moeda(precos[posicao])
                    cor = (232, 245, 233) if precos[posicao] <= media + 1e-9 else (255, 243, 224)
                elif posicao == 0 and not precos:
                    texto, alinhamento = STATUS_TEXTO.get(r["status"], ""), "C"
                else:
                    texto = "-"
            elif indice == 4 + max_precos:
                texto, fonte_negrito = (moeda(media) if media is not None else "-"), True
                cor = (255, 244, 204)
            elif indice == 5 + max_precos:
                texto = str(len(precos))
            else:
                texto = f"{r['stats']['cv']:.1f}%".replace(".", ",") if r["stats"] else "-"
            pdf.set_fill_color(*cor)
            pdf.set_draw_color(190, 190, 190)
            pdf.rect(x, y0, largura, altura, "DF")
            pdf.set_text_color(*((150, 30, 30) if (not precos and indice == 4) else (51, 51, 51)))
            pdf.set_font("Helvetica", "B" if fonte_negrito else "", 7.5 if fonte_negrito else 6.5)
            linhas_texto = texto if isinstance(texto, list) else [texto]
            deslocamento = (altura - 3.6 * len(linhas_texto)) / 2
            for n, linha in enumerate(linhas_texto):
                pdf.set_xy(x + 1, y0 + deslocamento + 3.6 * n)
                pdf.cell(largura - 2, 3.6, _seguro(linha), 0, 0, alinhamento)
            x += largura
        pdf.set_y(y0 + altura)
    pdf.ln(2)
    pdf.set_font("Helvetica", "I", 6.5)
    pdf.set_text_color(90, 90, 90)
    pdf.multi_cell(LARGURA, 3.5, new_x="LMARGIN", new_y="NEXT", text=_seguro(fonte or (
        f"Fonte: Compras.gov (Pesquisa de Preco - precos praticados), compras dos ultimos {JANELA_DIAS} dias, mesma unidade de fornecimento. "
        "Outliers e precos inexequiveis removidos. Detalhamento de cada item nas paginas seguintes.")))


def _pagina_totais(pdf: FPDF, resultados: list[dict]) -> None:
    """Valor estimado por item (média unitária x quantidade pedida) e total do orçamento. Só quando a lista trouxe quantidades."""
    pdf.add_page()
    pdf.set_y(35)
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_text_color(*AZUL)
    pdf.cell(0, 6, "VALOR ESTIMADO DO ORCAMENTO - media unitaria x quantidade pedida", ln=True)
    pdf.set_font("Helvetica", "", 6.5)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 4, "Itens sem quantidade informada ou sem precos validos nao entram no total.", ln=True)
    pdf.ln(1)
    colunas = [("#", 10), ("Item", 135), ("Qtd.", 22), ("Un.", 22), ("Media unitaria (R$)", 46), ("Valor total (R$)", 46)]
    _cabecalho_tabela(pdf, colunas)
    for numero, r in enumerate(resultados, start=1):
        media = r["stats"]["media"] if r["stats"] else None
        total = valor_total_item(r)
        desc = _quebrar(pdf, r["descricao"], colunas[1][1] - 2, 2)
        altura = max(6.5, 3.6 * len(desc) + 2)
        if pdf.get_y() + altura > pdf.h - 24:
            pdf.add_page()
            pdf.set_y(35)
            _cabecalho_tabela(pdf, colunas)
        x, y0 = pdf.l_margin, pdf.get_y()
        valores = [str(numero), desc, formatar_quantidade(r["quantidade_pedida"]) if r.get("quantidade_pedida") else "-", r.get("unidade_pedida") or "-",
                   moeda(media) if media is not None else "-", moeda(total) if total is not None else "-"]
        for indice, ((_, largura), valor) in enumerate(zip(colunas, valores)):
            pdf.set_fill_color(*((245, 248, 255) if numero % 2 == 0 else (255, 255, 255)))
            pdf.set_draw_color(190, 190, 190)
            pdf.rect(x, y0, largura, altura, "DF")
            pdf.set_text_color(51, 51, 51)
            pdf.set_font("Helvetica", "B" if indice == 5 else "", 7)
            linhas_texto = valor if isinstance(valor, list) else [valor]
            deslocamento = (altura - 3.6 * len(linhas_texto)) / 2
            for n, linha in enumerate(linhas_texto):
                pdf.set_xy(x + 1, y0 + deslocamento + 3.6 * n)
                pdf.cell(largura - 2, 3.6, _seguro(linha), 0, 0, "L" if indice == 1 else "C")
            x += largura
        pdf.set_y(y0 + altura)
    soma = sum(w for _, w in colunas[:5])
    pdf.set_fill_color(255, 244, 204)
    pdf.set_draw_color(190, 190, 190)
    pdf.set_font("Helvetica", "B", 8.5)
    pdf.set_text_color(*AZUL)
    pdf.cell(soma, 8, "VALOR TOTAL ESTIMADO DO ORCAMENTO (R$)  ", 1, 0, "R", True)
    pdf.cell(colunas[5][1], 8, moeda(valor_total_orcamento(resultados)), 1, 1, "C", True)


COLUNAS_PRECOS = [("ID Compra", 26), ("Data", 16), ("UASG", 15), ("Unid.", 22), ("Qtd", 12), ("V. Unitario", 20), ("CNPJ", 28), ("Fornecedor", 62), ("UF", 8), ("Codigo", 16), ("Orgao", 56)]


def _pagina_item(pdf: FPDF, numero: int, r: dict) -> None:
    pdf.add_page()
    pdf.set_y(35)
    pdf.set_fill_color(*AZUL_CARTAO)
    pdf.set_text_color(*DOURADO)
    pdf.set_font("Helvetica", "B", 8)
    tipo_item = "SERVICO" if r.get("tipo") == "Serviço" else "MATERIAL"
    pdf.cell(LARGURA, 6, f"  ITEM {numero} ({tipo_item}) - {STATUS_TEXTO.get(r['status'], '').upper()}", 0, 1, "L", True)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(*AZUL_CARTAO)
    for linha in _quebrar(pdf, r["descricao"], LARGURA - 4, 3):
        pdf.cell(LARGURA, 5, "  " + linha, 0, 1, "L", True)
    pdf.ln(3)

    if r["catmats"]:
        pdf.set_font("Helvetica", "B", 7)
        pdf.set_text_color(*AZUL)
        rotulo_codigo = "CATSERV" if r.get("tipo") == "Serviço" else "CATMAT"
        criterio = {"estimativa": "escolhidos pela estimativa de preco informada", "manual": "escolhido pelo usuario"}.get(r.get("escolha"), f">= {r.get('limiar', LIMIAR_CORRESPONDENCIA):.0f}%")
        pdf.cell(0, 5, f"{rotulo_codigo} correspondentes ({criterio}) - confira a descricao antes de usar no processo", ln=True)
        _cabecalho_tabela(pdf, [(rotulo_codigo, 20), ("Corresp.", 18), ("Registros", 18), ("Descricao no catalogo", LARGURA - 56)], 6)
        pdf.set_font("Helvetica", "", 6.5)
        pdf.set_text_color(51, 51, 51)
        for k in r["catmats"]:
            pdf.cell(20, 5.5, str(k["codigo"]), 1, 0, "C")
            pdf.cell(18, 5.5, f"{k['correspondencia']:.0f}%", 1, 0, "C")
            pdf.cell(18, 5.5, str(k["registros"]), 1, 0, "C")
            pdf.cell(LARGURA - 56, 5.5, _truncar(pdf, k["descricao"], LARGURA - 58), 1, 1, "L")
        pdf.ln(3)

    if not r["precos"]:
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(150, 30, 30)
        if r["status"] == "sem_catmat":
            texto = f"Nenhum item CATMAT/CATSERV com {r.get('limiar', LIMIAR_CORRESPONDENCIA):.0f}% ou mais de correspondencia e precos praticados nos ultimos {JANELA_DIAS} dias."
            if r.get("proximos"):
                texto += " Codigos mais proximos (nao atingiram o minimo" + (" ou nao tem precos no periodo): " if r.get("tipo") == "Serviço" else "): ")
                texto += " | ".join(f"{p['codigo']} ({p['correspondencia']:.0f}%) {p['descricao'][:70]}" for p in r["proximos"])
        else:
            texto = "Os precos encontrados nao formaram um conjunto coerente (a ate 30% da media) apos a remocao de outliers."
        pdf.multi_cell(LARGURA, 5, _seguro(texto), new_x="LMARGIN", new_y="NEXT")
        return

    s = r["stats"]
    _cartoes(pdf, [("MINIMO", f"R$ {moeda(s['min'])}"), ("MEDIA UNITARIA", f"R$ {moeda(s['media'])}"), ("MEDIANA", f"R$ {moeda(s['mediana'])}"),
                   ("MAXIMO", f"R$ {moeda(s['max'])}"), ("DESVIO PADRAO", f"R$ {moeda(s['desvio'])}"), ("CV", f"{s['cv']:.2f}%".replace(".", ","))], pdf.get_y())
    _cabecalho_tabela(pdf, COLUNAS_PRECOS)
    pdf.set_font("Helvetica", "", 6)
    pdf.set_text_color(51, 51, 51)
    for indice, p in enumerate(r["precos"]):
        pdf.set_fill_color(245, 248, 255)
        valores = [p["id_compra"], str(p["data"] or "")[:10], p["uasg"], p["sigla"], f"{p['quantidade']:g}" if p["quantidade"] is not None else "",
                   f"R$ {moeda(p['preco'])}", p["cnpj"], p["fornecedor"], p["uf"], p["catmat"], p["nome_uasg"]]
        for (_, largura), valor in zip(COLUNAS_PRECOS, valores):
            pdf.cell(largura, 6, _truncar(pdf, valor, largura - 2), 1, 0, "C", indice % 2 == 1)
        pdf.ln()
    pdf.ln(3)
    pdf.set_font("Helvetica", "", 7)
    pdf.set_text_color(70, 70, 70)
    faixa = f"R$ {moeda(s['media'] * (1 - TOLERANCIA))} a R$ {moeda(s['media'] * (1 + TOLERANCIA))}"
    validos = r["faixa_validos"]
    notas = [
        f"Unidade de fornecimento considerada: {r['unidade']}.",
        f"Universo analisado: {r.get('universo', r['brutos'])} registros; outliers/precos inexequiveis removidos: {r['outliers']}"
        + (f" (faixa valida: R$ {moeda(validos[0])} a R$ {moeda(validos[1])})." if validos else "."),
        f"Os {len(r['precos'])} precos listados ficam a ate {TOLERANCIA * 100:.0f}% da media (faixa de R$ {faixa.replace('R$ ', '', 1)}).",
    ]
    if r.get("faixa_estimativa"):
        baixo, alto = r["faixa_estimativa"]
        notas.append(f"Estimativa de preco informada pelo requisitante: R$ {moeda(r['estimativa'])} por unidade. Os precos listados foram escolhidos entre as compras de "
                     f"R$ {moeda(baixo)} a R$ {moeda(alto)} (+-40% da estimativa); compras fora dessa faixa nao entraram na selecao.")
    if r["status"] == "insuficiente":
        notas.append(f"ATENCAO: menos de {MIN_PRECOS} precos - complementar com outras fontes (IN 65/2021).")
    for nota in notas:
        pdf.multi_cell(LARGURA, 4, _seguro(nota), new_x="LMARGIN", new_y="NEXT")


def gerar_pdf(
    resultados: list[dict],
    titulo: str = "Cotacao Rapida - Mapa Comparativo de Precos",
    pagina_item=None,
    textos: tuple[str, ...] | None = None,
    argumentos_mapa: dict | None = None,
    anexos=None,
) -> bytes:
    """PDF no layout da Cotação: mapa comparativo na 1ª página, uma página por item e textos finais.

    Os parâmetros permitem reaproveitar o layout em outros relatórios (ex.: pesquisa de notas fiscais em lote)."""
    pdf = PDFCotacaoRapida(orientation="L", unit="mm", format="A4", titulo=titulo)
    pdf.set_margins(8, 35, 8)
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    _pagina_mapa(pdf, resultados, **(argumentos_mapa or {}))
    if tem_quantidades(resultados):
        _pagina_totais(pdf, resultados)
    for numero, resultado in enumerate(resultados, start=1):
        (pagina_item or _pagina_item)(pdf, numero, resultado)
    if anexos:
        anexos(pdf)  # páginas extras entre o detalhe dos itens e os textos finais
    pdf.add_page()
    pdf.ln(5)
    for texto in textos or (JUSTIFICATIVA_COTACAO, metodologia()):
        titulo_texto, *paragrafos = texto.split("\n\n")
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(*AZUL)
        pdf.multi_cell(0, 8, _seguro(titulo_texto), align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
        pdf.set_font("Helvetica", "", 9.5)
        pdf.set_text_color(51, 51, 51)
        for paragrafo in paragrafos:
            pdf.multi_cell(0, 5.5, _seguro(paragrafo), align="J", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2.5)
        pdf.ln(4)
    return bytes(pdf.output())


COL_CATMAT = "CATMAT"
COL_DESC_CATMAT = "Descrição do item no CATMAT"
COL_TOTAL = "Valor total (média × qtd.)"


def _catmat_da_linha(r: dict) -> tuple[str, float | None, str, str]:
    """(código, % de casamento, descrição no catálogo, outros códigos) do CATMAT/CATSERV mais próximo do item."""
    if r["catmats"]:
        primeiro = r["catmats"][0]
        return str(primeiro["codigo"]), round(primeiro["correspondencia"], 1), str(primeiro["descricao"]), " | ".join(f"{k['codigo']} ({k['correspondencia']:.0f}%)" for k in r["catmats"][1:])
    proximo = r.get("melhor_proximo")
    if proximo:
        return "-", None, f"Nenhum com preços no período; mais próximo: {proximo['codigo']} ({proximo['correspondencia']:.0f}%) {proximo['descricao']}", ""
    return "-", None, "Nenhuma correspondência no catálogo", ""


def tabela_mapa(resultados: list[dict]) -> pd.DataFrame:
    """Mapa comparativo em DataFrame (usado na tela e no Excel), com as mesmas colunas da tabela de itens da Cotação Direta.
    Com quantidades na lista, traz Qtd., Un. e o valor total de cada item (média unitária x quantidade)."""
    com_quantidade = tem_quantidades(resultados)
    com_estimativa = any(r.get("estimativa") for r in resultados)
    material_e_servico = len({r.get("tipo") == "Serviço" for r in resultados}) > 1 or any(r.get("tipo") == "Serviço" for r in resultados)
    rotulo_codigo = "CATMAT/CATSERV" if material_e_servico else COL_CATMAT
    linhas = []
    for numero, r in enumerate(resultados, start=1):
        precos = [p["preco"] for p in r["precos"]]
        codigo, casamento, descricao_catalogo, outros = _catmat_da_linha(r)
        linha: dict = {"Item": numero, "Descrição do item": r["descricao"]}
        if com_quantidade:
            linha["Qtd."] = r.get("quantidade_pedida")
            linha["Un."] = r.get("unidade_pedida") if r.get("quantidade_pedida") else None
        linha.update({"Tipo": "Serviço" if r.get("tipo") == "Serviço" else "Material", rotulo_codigo: codigo, "% casamento": casamento,
                      COL_DESC_CATMAT: descricao_catalogo, "Outros códigos": outros or None, "Un. de fornecimento": r["unidade"] or "-"})
        for n in range(MAX_PRECOS):
            linha[f"Preço {n + 1}"] = precos[n] if n < len(precos) else None
        s = r["stats"]
        linha["Média unitária"] = s["media"] if s else None
        if com_quantidade:
            linha[COL_TOTAL] = valor_total_item(r)
        if com_estimativa:
            linha["Estimativa (R$)"] = r.get("estimativa")
            linha["CATMAT escolhido por"] = {"estimativa": "Estimativa de preço", "manual": "Escolha do usuário"}.get(r.get("escolha"), "Mais parecido") if r.get("estimativa") else None
            linha["Conferência com a estimativa"] = r.get("texto_validacao") or None
        linha.update({
            "Mediana": s["mediana"] if s else None,
            "Mínimo": s["min"] if s else None,
            "Máximo": s["max"] if s else None,
            "Desvio padrão": s["desvio"] if s else None,
            "CV (%)": s["cv"] if s else None,
            "Nº de preços": len(precos),
            "Situação": STATUS_TEXTO.get(r["status"], ""),
        })
        linhas.append(linha)
    return pd.DataFrame(linhas)


def aplicar_formulas_mapa(planilha, com_quantidade: bool, coluna_descricao: str = "Descrição do item") -> None:
    """Na aba do mapa: média unitária e valor total viram fórmulas (mudou um preço ou uma quantidade, recalcula) e entra a linha de total do orçamento."""
    cabecalho = {str(c.value): c.column for c in planilha[1] if c.value}
    ultima = planilha.max_row
    letras = {nome: get_column_letter(coluna) for nome, coluna in cabecalho.items()}
    precos = [letras[n] for n in cabecalho if n.startswith("Preço ")]
    if "Média unitária" in letras and precos:
        for linha in range(2, ultima + 1):
            planilha[f"{letras['Média unitária']}{linha}"] = f'=IFERROR(AVERAGE({precos[0]}{linha}:{precos[-1]}{linha}),"")'
    if com_quantidade and COL_TOTAL in letras and "Qtd." in letras and "Média unitária" in letras:
        for linha in range(2, ultima + 1):
            media, qtd = f"{letras['Média unitária']}{linha}", f"{letras['Qtd.']}{linha}"
            planilha[f"{letras[COL_TOTAL]}{linha}"] = f'=IF(AND(ISNUMBER({media}),ISNUMBER({qtd})),ROUND({media}*{qtd},2),"")'
        total = ultima + 1
        planilha[f"{letras[coluna_descricao]}{total}"] = "VALOR TOTAL DO ORÇAMENTO"
        planilha[f"{letras[COL_TOTAL]}{total}"] = f"=SUM({letras[COL_TOTAL]}2:{letras[COL_TOTAL]}{ultima})"
        for nome in (coluna_descricao, COL_TOTAL):
            celula = planilha[f"{letras[nome]}{total}"]
            celula.font = Font(bold=True, color="001A4D")
            celula.fill = PatternFill("solid", fgColor="FFF4CC")
            celula.number_format = '"R$" #,##0.00'


def gerar_excel(resultados: list[dict]) -> bytes:
    detalhes = []
    for numero, r in enumerate(resultados, start=1):
        for p in r["precos"]:
            detalhes.append({
                "Item": numero, "Descrição pesquisada": r["descricao"], "Tipo": "Serviço" if r.get("tipo") == "Serviço" else "Material", "Código": p["catmat"], "Descrição no catálogo": p["descricao"],
                "Unidade": r["unidade"], "ID Compra": p["id_compra"], "Data": str(p["data"] or "")[:10], "UASG": p["uasg"], "Órgão": p["nome_uasg"],
                "UF": p["uf"], "Fornecedor": p["fornecedor"], "CNPJ": p["cnpj"], "Quantidade": p["quantidade"], "Valor unitário": p["preco"],
            })
    saida = io.BytesIO()
    with pd.ExcelWriter(saida, engine="openpyxl") as escritor:
        tabela_mapa(resultados).to_excel(escritor, index=False, sheet_name="Mapa Comparativo")
        pd.DataFrame(detalhes).to_excel(escritor, index=False, sheet_name="Preços selecionados")
        for planilha in escritor.book.worksheets:
            for celula in planilha[1]:
                celula.font = Font(color="FFFFFF", bold=True)
                celula.fill = PatternFill("solid", fgColor="001A4D")
                celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for coluna in planilha.columns:
                maior = max(len(str(c.value or "")) for c in coluna)
                planilha.column_dimensions[get_column_letter(coluna[0].column)].width = min(max(maior + 2, 11), 60)
                if coluna[0].value and ("Preço" in str(coluna[0].value) or "Média" in str(coluna[0].value) or "Valor" in str(coluna[0].value)
                                        or coluna[0].value in ("Mediana", "Mínimo", "Máximo", "Desvio padrão")):
                    for celula in coluna[1:]:
                        celula.number_format = '"R$" #,##0.00'
            planilha.freeze_panes = "A2"
        aplicar_formulas_mapa(escritor.book["Mapa Comparativo"], tem_quantidades(resultados))
    return saida.getvalue()
